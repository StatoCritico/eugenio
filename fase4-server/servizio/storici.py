"""Dati storici scaricabili (pagina "Dati storici" del sito): catalogo, stima ed estrazione in CSV o Excel.

Rotte (da includere in app.py con `app.include_router(storici.router)`):
  GET  /api/storici/catalogo                  dataset presenti sul server, grandezze, periodi, limiti
  POST /api/storici/stima                     righe e dimensione del file per una selezione
  POST /api/storici/estrai?formato=csv|xlsx   il file (CSV in streaming; Excel con foglio "Info")

Si possono scaricare SOLO dati che le licenze permettono di ridistribuire, sempre con attribuzione:
  * meteo orario (ICON storico) e giornaliero (ERA5 corretto) e i loro derivati (gradi giorno, dati
    climatici per la Legge 10): Open-Meteo.com CC BY 4.0, ERA5 Copernicus, DWD ICON; correzione EuGEnio;
  * elenco dei comuni con il punto meteo: ISTAT CC BY 4.0;
  * prezzi Eurostat dell'energia per le PMI: Eurostat, riuso consentito con citazione della fonte;
  * gradi giorno, zona climatica e altitudine dell'Allegato A del DPR 412/93: atto normativo, libero
    (art. 5 L. 633/1941). Il server legge `DATI/dpr412.json`, copia del file `sito/dati/dpr412.json`
    prodotto da fase5-sito/costruisci_dpr412.py (lo stesso file che usa la pagina Legge 10).
NON sono scaricabili, e ogni richiesta che li nomina è rifiutata con 403: prezzi GME (PUN, zonali, PSV, IGI),
quotazioni EEX/TTF e future, PFOR, PSV ARERA, dati Zornade, indirizzario ANNCSU grezzo. Il modulo legge
solo percorsi fissi della cartella dati (mai `mercati/` fuori da `mercati/eurostat/`).

Nessun dato degli utenti viene salvato: le richieste contengono solo i filtri; il conteggio per il limite
d'uso per IP sta in memoria e si azzera al riavvio. I file Excel passano da un file temporaneo cancellato
subito dopo l'invio.
"""
from __future__ import annotations

import datetime as dt
import io
import math
import os
import re
import tempfile
import threading
import time
from collections import defaultdict, deque
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path
from typing import Iterator, Literal
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from . import config as C
from .report import CITAZIONE, LICENZA_DATI, LICENZA_PAGINA

router = APIRouter(prefix="/api/storici", tags=["dati storici"])
ROMA = ZoneInfo("Europe/Rome")
INIZIO_ICON = dt.date(2022, 11, 14)

# ---------------------------------------------------------------------------------------------- licenze
_ANNO = dt.date.today().year
METEO_FONTE = ("Weather data by Open-Meteo.com (CC BY 4.0). ERA5: Copernicus Climate Change Service (C3S), "
               "ECMWF – contiene informazioni modificate del Copernicus Climate Change Service. "
               "ICON: Deutscher Wetterdienst (DWD). Correzione ERA5 → ICON ed elaborazioni: EuGEnio (Stato Critico).")
ISTAT_FONTE = "Elenco dei comuni, codici e confini: ISTAT (CC BY 4.0); punto meteo: elaborazione EuGEnio sulle sezioni di censimento 2021."
EUROSTAT_FONTE = "Fonte: Eurostat (nrg_pc_205 elettricità, nrg_pc_203 gas), dati forniti da ISTAT/ARERA."
LIC_CCBY = ("CC BY 4.0", "https://creativecommons.org/licenses/by/4.0/deed.it")
DPR412_URL = ("https://www.normattiva.it/uri-res/N2Ls?urn:nir:stato:decreto.del.presidente.della.repubblica:"
              "1993-08-26;412!vig=")
DPR412_FONTE = ("DPR 26 agosto 1993, n. 412, Allegato A (testo vigente su Normattiva); abbinamento ai codici ISTAT "
                "attuali: EuGEnio (Stato Critico) su elenchi ISTAT (CC BY 4.0).")
LIC_DPR412 = ("Atto normativo, non coperto da diritto d'autore (art. 5 L. 633/1941); abbinamento ISTAT CC BY 4.0",
              DPR412_URL)
LIC_EUROSTAT = ("Riuso consentito con citazione della fonte (Commissione europea, decisione 2011/833/UE)",
                "https://ec.europa.eu/eurostat/about-us/policies/copyright")

# Licenza delle elaborazioni e raccolte di EuGEnio (stessa formula di citazione dei report e del sito)
RIUSO = (f"Elaborazioni e raccolte di EuGEnio: {LICENZA_DATI[0]} ({LICENZA_DATI[1]}). Si possono riutilizzare, anche "
         "modificate, citando EuGEnio come fonte nei documenti che le usano e mantenendo le fonti indicate sopra, "
         "che restano soggette alle loro licenze (i testi normativi sono atti pubblici). "
         f"Dettagli: {LICENZA_PAGINA}")

# Dati interni: mai scaricabili. Il controllo è sui nomi richiesti (dataset, grandezze, tipologie, serie).
VIETATI = ("pun", "zonal", "mgp", "psv", "igi", "ttf", "eex", "future", "forward", "pfor", "arera", "cmem",
           "gme", "zornade", "anncsu", "indirizz", "civic", "catast", "particell", "indici_importati")
NON_SCARICABILI = [
    {"nome": "Prezzi GME (PUN, prezzi zonali, MGP-GAS, PSV, IGI)", "motivo": "condizioni d'uso del GME: solo uso interno"},
    {"nome": "Quotazioni EEX/TTF e future trimestrali", "motivo": "dati di borsa EEX: solo uso interno (calcolo del PFOR)"},
    {"nome": "PFOR,t ricostruito e PSV ARERA", "motivo": "derivati da dati non ridistribuibili"},
    {"nome": "Dati Zornade (catasto, geocodifica)", "motivo": "condizioni del servizio Zornade"},
    {"nome": "Indirizzario ANNCSU grezzo", "motivo": "si usa solo per geolocalizzare gli edifici"},
]
MSG_VIETATO = ("«{x}» non è scaricabile. I prezzi GME (PUN, zonali, PSV, IGI), le quotazioni EEX/TTF e i future, "
               "il PFOR, il PSV ARERA, i dati Zornade e l'indirizzario ANNCSU servono solo per elaborazioni interne: "
               "le loro condizioni d'uso non ne permettono la ridistribuzione.")


def _vietato(x: str | None) -> bool:
    s = (x or "").lower()
    return any(v in s for v in VIETATI)


# ---------------------------------------------------------------------------------------------- grandezze
@dataclass(frozen=True)
class Uscita:
    """Una colonna del file: nome, unità, descrizione, funzione di aggregazione, decimali."""
    nome: str
    unita: str
    descrizione: str
    sorgente: str = ""
    agg: str = ""            # mean | sum | min | max ("" = nessuna aggregazione)
    scala: float = 1.0
    dec: int = 2
    larghezza: int = 5       # caratteri medi nel CSV, per la stima della dimensione


@dataclass(frozen=True)
class Grandezza:
    id: str
    nome: str
    gruppo: str
    base: tuple[Uscita, ...]          # colonne alla frequenza nativa
    aggregate: tuple[Uscita, ...]     # colonne alle frequenze aggregate
    predefinita: bool = True

    def catalogo(self, nativa: str) -> dict:
        u = self.base[0]
        a = self.aggregate[0] if self.aggregate else None
        aggr = {"mean": "media", "sum": "somma", "min": "minimo", "max": "massimo"}
        return {"id": self.id, "nome": self.nome, "gruppo": self.gruppo, "unita": u.unita,
                "unita_aggregata": a.unita if a else None,
                "aggregazione": "media, minimo e massimo" if len(self.aggregate) > 1 else aggr.get(a.agg) if a else None,
                "predefinita": self.predefinita}


def _g(id, nome, gruppo, col, unita, descr, agg, unita_agg=None, scala=1.0, dec=2, dec_agg=None, nome_agg=None,
       predefinita=True):
    base = (Uscita(col, unita, descr, id, "", 1.0, dec),)
    ag = (Uscita(nome_agg or col, unita_agg or unita, descr, id, agg, scala, dec if dec_agg is None else dec_agg),)
    return Grandezza(id, nome, gruppo, base, ag, predefinita)


ORARIE = [
    Grandezza("temperature_2m", "Temperatura dell'aria a 2 m", "Temperature",
              (Uscita("temperatura", "°C", "temperatura dell'aria a 2 m", "temperature_2m", "", 1, 2),),
              (Uscita("t_media", "°C", "media delle temperature orarie", "temperature_2m", "mean", 1, 2),
               Uscita("t_min", "°C", "minima oraria del periodo", "temperature_2m", "min", 1, 2),
               Uscita("t_max", "°C", "massima oraria del periodo", "temperature_2m", "max", 1, 2))),
    _g("relative_humidity_2m", "Umidità relativa", "Umidità", "umidita_relativa", "%", "umidità relativa a 2 m",
       "mean", dec=1, nome_agg="umidita_relativa_media"),
    _g("shortwave_radiation", "Radiazione solare globale orizzontale", "Radiazione", "radiazione_globale", "W/m²",
       "irradianza globale su piano orizzontale, media dell'ora precedente", "sum", "kWh/m²", 1 / 1000, 1, 3),
    _g("diffuse_radiation", "Radiazione solare diffusa orizzontale", "Radiazione", "radiazione_diffusa", "W/m²",
       "irradianza diffusa su piano orizzontale, media dell'ora precedente", "sum", "kWh/m²", 1 / 1000, 1, 3),
    _g("wind_speed_10m", "Velocità del vento a 10 m", "Vento", "vento", "km/h", "velocità del vento a 10 m", "mean",
       dec=1, nome_agg="vento_medio"),
    _g("cloud_cover", "Nuvolosità", "Nuvolosità", "nuvolosita", "%", "copertura nuvolosa totale", "mean", dec=0,
       dec_agg=1, nome_agg="nuvolosita_media"),
    _g("precipitation", "Precipitazione (pioggia e neve fusa)", "Precipitazioni", "precipitazione", "mm",
       "precipitazione totale", "sum", dec=1),
    _g("snowfall", "Neve fresca", "Precipitazioni", "neve", "cm", "neve fresca caduta", "sum", dec=1),
]
GIORNALIERE = [
    _g("t_media", "Temperatura media", "Temperature", "t_media", "°C", "temperatura media giornaliera", "mean"),
    _g("t_min", "Temperatura minima", "Temperature", "t_min", "°C", "temperatura minima (del giorno o del periodo)",
       "min"),
    _g("t_max", "Temperatura massima", "Temperature", "t_max", "°C", "temperatura massima (del giorno o del periodo)",
       "max"),
    _g("ur_media", "Umidità relativa media", "Umidità", "ur_media", "%", "umidità relativa media", "mean", dec=1),
    _g("rad_globale_kwh_m2", "Radiazione solare globale orizzontale", "Radiazione", "rad_globale", "kWh/m²",
       "irradiazione globale su piano orizzontale", "sum", dec=3),
    _g("rad_diffusa_kwh_m2", "Radiazione solare diffusa orizzontale", "Radiazione", "rad_diffusa", "kWh/m²",
       "irradiazione diffusa su piano orizzontale", "sum", dec=3),
    _g("vento_medio_kmh", "Velocità media del vento a 10 m", "Vento", "vento_medio", "km/h",
       "velocità media del vento a 10 m", "mean", dec=1),
    _g("nuvolosita_media", "Nuvolosità media", "Nuvolosità", "nuvolosita_media", "%", "copertura nuvolosa media",
       "mean", dec=1),
    _g("pioggia_mm", "Precipitazione (pioggia e neve fusa)", "Precipitazioni", "precipitazione", "mm",
       "precipitazione totale", "sum", dec=1),
    _g("neve_cm", "Neve fresca", "Precipitazioni", "neve", "cm", "neve fresca caduta", "sum", dec=1),
    _g("gg20", "Gradi giorno di riscaldamento (base 20 °C)", "Gradi giorno", "gg20", "°C·giorno",
       "somma di (20 − t media) sui giorni con t media < 20 °C", "sum", dec=1, predefinita=False),
    _g("gc24", "Gradi giorno di raffrescamento (base 24 °C)", "Gradi giorno", "gc24", "°C·giorno",
       "somma di (t media − 24) sui giorni con t media > 24 °C", "sum", dec=1, predefinita=False),
]
_GG = {g.id: g for g in GIORNALIERE}
GRADI_GIORNO = [
    Grandezza("gg20", "Gradi giorno di riscaldamento (base 20 °C)", "Gradi giorno", _GG["gg20"].base, _GG["gg20"].aggregate),
    Grandezza("gc24", "Gradi giorno di raffrescamento (base 24 °C)", "Gradi giorno", _GG["gc24"].base, _GG["gc24"].aggregate),
    Grandezza("t_media", "Temperatura media (per confronto)", "Temperature", _GG["t_media"].base,
              _GG["t_media"].aggregate, predefinita=False),
]

FREQUENZE = {"oraria": "oraria", "giornaliera": "giornaliera", "mensile": "mensile", "annuale": "annuale",
             "stagione": "per stagione di riscaldamento (1 luglio – 30 giugno)"}

LEGGE10_TIPOLOGIE = {
    "mensili": "Dati mensili come nella UNI 10349-1 (12 righe per comune)",
    "riepilogo": "Riepilogo per comune: gradi giorno, zona climatica, temperature di progetto",
    "stagioni": "Gradi giorno stagione per stagione (periodo di accensione della zona)",
}
TASSAZIONE = {"X_TAX": "senza imposte e oneri recuperabili", "X_VAT": "IVA e imposte recuperabili escluse",
              "I_TAX": "tutte le imposte incluse"}
SERIE_EUROSTAT = {"elettricita_pmi": "Elettricità, clienti non domestici (nrg_pc_205)",
                  "gas_pmi": "Gas naturale, clienti non domestici (nrg_pc_203)"}

DATASET = {
    "meteo_orario": {
        "titolo": "Meteo orario (ICON storico)",
        "descrizione": "Le otto grandezze ora per ora dal 14 novembre 2022, dalle previsioni storiche del modello "
                       "ICON (DWD) archiviate da Open-Meteo: lo stesso riferimento delle previsioni di EuGEnio. "
                       "Ore in UTC e in ora italiana; le aggregazioni sono per giorno civile italiano.",
        "fonte": METEO_FONTE, "licenza": LIC_CCBY, "periodo": "date", "area": True,
        "frequenze": ["oraria", "giornaliera", "mensile", "annuale", "stagione"], "nativa": "oraria",
    },
    "meteo_giornaliero": {
        "titolo": "Meteo giornaliero (ERA5 corretto)",
        "descrizione": "Valori giornalieri dal 2015 dalla rianalisi ERA5 (Copernicus), corretta comune per comune "
                       "verso ICON sul periodo in cui ci sono entrambi: una serie lunga e omogenea per confrontare "
                       "gli anni e normalizzare i consumi.",
        "fonte": METEO_FONTE, "licenza": LIC_CCBY, "periodo": "date", "area": True,
        "frequenze": ["giornaliera", "mensile", "annuale", "stagione"], "nativa": "giornaliera",
    },
    "gradi_giorno": {
        "titolo": "Gradi giorno",
        "descrizione": "Gradi giorno di riscaldamento (base 20 °C) e di raffrescamento (base 24 °C) dalla temperatura "
                       "media giornaliera ERA5 corretta, sull'anno intero (non solo sul periodo di accensione "
                       "del DPR 412/93: per quello usa il dataset della Legge 10).",
        "fonte": METEO_FONTE, "licenza": LIC_CCBY, "periodo": "date", "area": True,
        "frequenze": ["giornaliera", "mensile", "annuale", "stagione"], "nativa": "giornaliera",
    },
    "legge10": {
        "titolo": "Dati climatici per la Relazione Legge 10",
        "descrizione": "Per ogni comune, dalle ultime stagioni complete: dati mensili come nella UNI 10349-1, gradi "
                       "giorno sul periodo di accensione, zona climatica risultante, temperature di progetto stimate. "
                       "Per la relazione tecnica restano obbligatori i valori normativi (UNI 10349, DPR 412/93).",
        "fonte": METEO_FONTE, "licenza": LIC_CCBY, "periodo": "stagioni", "area": True, "frequenze": [],
        "avvertenza": "Valori dal clima osservato (ERA5 corretto). Per la relazione tecnica ex Legge 10 restano "
                      "obbligatori i valori normativi (UNI 10349, DPR 412/93).",
    },
    "dpr412": {
        "titolo": "Gradi giorno e zone del DPR 412/93",
        "descrizione": "Valori di legge dell'Allegato A del DPR 412/93 per comune: gradi giorno, zona climatica e "
                       "altitudine della casa comunale, abbinati al codice ISTAT attuale. Per i comuni nati da "
                       "fusione dopo il 1993 ci sono più righe, una per ogni comune d'origine.",
        "fonte": DPR412_FONTE, "licenza": LIC_DPR412, "periodo": None, "area": True, "frequenze": [],
        "avvertenza": "Sono i valori normativi da usare nella relazione tecnica (insieme alla UNI 10349); "
                      "per i comuni nati da fusione il DPR non ha un valore unico.",
    },
    "comuni": {
        "titolo": "Elenco dei comuni con il punto meteo",
        "descrizione": "Codice ISTAT e catastale, provincia, regione, capoluogo e coordinate del punto usato per il "
                       "meteo (baricentro degli edifici del censimento 2021 o, in mancanza, del territorio comunale).",
        "fonte": ISTAT_FONTE, "licenza": LIC_CCBY, "periodo": None, "area": True, "frequenze": [],
    },
    "eurostat_pmi": {
        "titolo": "Prezzi dell'energia per le PMI (Eurostat)",
        "descrizione": "Prezzi semestrali di elettricità e gas per i clienti non domestici in Italia, per fascia di "
                       "consumo e livello di tassazione. Valori nell'unità di Eurostat (di solito €/kWh).",
        "fonte": EUROSTAT_FONTE, "licenza": LIC_EUROSTAT, "periodo": "semestri", "area": False, "frequenze": [],
    },
}


# ---------------------------------------------------------------------------------------------- file sul server
def _p_giornaliero() -> Path:
    return C.DATI / "giornaliero.parquet"


def _p_orario(istat: str) -> Path:
    return C.DATI / "orario" / f"{istat}.parquet"


def _p_dpr412() -> Path:
    return C.DATI / "dpr412.json"


@lru_cache(maxsize=2)
def _dpr412_cache(mtime: float) -> pd.DataFrame:
    """Il JSON compatto del sito in righe: una per comune, una per comune d'origine nelle fusioni."""
    import json
    colonne = ["istat", "gradi_giorno", "zona", "altitudine", "nome_nel_dpr", "tipo", "nota"]
    try:
        doc = json.loads(_p_dpr412().read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return pd.DataFrame(columns=colonne)
    righe = []
    for istat, v in (doc.get("comuni") or {}).items():
        if not isinstance(v, list) or len(v) < 5:
            continue
        gg, zona, alt, nome, nota = v[:5]
        origini = v[5] if len(v) > 5 and isinstance(v[5], list) else []
        if gg is None and origini:          # fusione: una riga per comune d'origine
            for o in origini:
                righe.append([istat, o[0], o[1], o[2], o[3], "fusione", nota])
        elif gg is None:
            righe.append([istat, None, None, None, "", "assente", nota])
        else:
            righe.append([istat, gg, zona, alt, nome, "incorporazione" if origini else "diretto", nota])
    df = pd.DataFrame(righe, columns=colonne)
    for c in ("gradi_giorno", "altitudine"):     # interi, vuoti dove il DPR non ha valore
        df[c] = pd.to_numeric(df[c], errors="coerce").round().astype("Int64")
    return df


def dpr412() -> pd.DataFrame:
    return _dpr412_cache(_mtime(_p_dpr412()))


def _d_eurostat() -> Path:
    return C.DATI / "mercati" / "eurostat"


def _file_eurostat() -> list[Path]:
    d = _d_eurostat()
    return sorted(p for p in d.glob("*.parquet") if not _vietato(p.stem)) if d.is_dir() else []


def _mtime(p: Path) -> float:
    try:
        return p.stat().st_mtime
    except OSError:
        return 0.0


@lru_cache(maxsize=2)
def _comuni_cache(mtime: float, regioni: tuple, solo: tuple) -> pd.DataFrame:
    p = C.DATI / "comuni.parquet"
    if p.exists():
        df = pd.read_parquet(p)
    elif (C.DATI / "comuni.json").exists():
        df = pd.read_json(C.DATI / "comuni.json", dtype={"istat": str, "regione_cod": str})
    else:
        return pd.DataFrame(columns=["istat", "nome", "provincia", "regione_cod", "regione", "capoluogo"])
    df["istat"] = df.istat.astype(str).str.zfill(6)
    if "regione_cod" not in df:
        df["regione_cod"] = ""
    df["regione_cod"] = df.regione_cod.astype(str).str.zfill(2)
    if "capoluogo" not in df:
        df["capoluogo"] = False
    df["capoluogo"] = df.capoluogo.fillna(False).astype(bool)
    if "regione" not in df:
        from meteo_storico.comuni import REGIONI as NOMI
        df["regione"] = df.regione_cod.map(NOMI)
    # stesso filtro di meteo.comuni(): solo i comuni di cui il server tiene i dati
    if solo:
        df = df[df.istat.isin(solo)]
    elif regioni:
        df = df[df.regione_cod.isin(regioni)]
    return df.sort_values("istat").reset_index(drop=True)


def comuni() -> pd.DataFrame:
    p = C.DATI / "comuni.parquet"
    return _comuni_cache(_mtime(p) or _mtime(C.DATI / "comuni.json"), tuple(C.REGIONI), tuple(C.SOLO_COMUNI))


def _estremi(p: Path, col: str) -> tuple:
    """Minimo e massimo di una colonna dalle statistiche del parquet (senza leggere i dati)."""
    try:
        f = pq.ParquetFile(p)
        i = f.schema_arrow.get_field_index(col)
        lo, hi = [], []
        for rg in range(f.metadata.num_row_groups):
            s = f.metadata.row_group(rg).column(i).statistics
            if s is None or not s.has_min_max:
                raise ValueError
            lo.append(s.min)
            hi.append(s.max)
        if lo:
            return min(lo), max(hi)
    except Exception:
        pass
    x = pd.read_parquet(p, columns=[col])[col]
    return (x.min(), x.max()) if len(x) else (None, None)


def _a_data(x) -> dt.date | None:
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return None
    t = pd.Timestamp(x)
    if t.tzinfo is not None:
        t = t.tz_convert(ROMA)
    return t.date()


def _semestre(lbl: str) -> tuple[dt.date, dt.date]:
    m = re.match(r"^(\d{4})(?:-S([12]))?$", str(lbl).strip())
    if not m:
        raise ValueError(lbl)
    a = int(m.group(1))
    if m.group(2) == "1":
        return dt.date(a, 1, 1), dt.date(a, 6, 30)
    if m.group(2) == "2":
        return dt.date(a, 7, 1), dt.date(a, 12, 31)
    return dt.date(a, 1, 1), dt.date(a, 12, 31)


@lru_cache(maxsize=4)
def _disponibilita(firma: tuple) -> dict:
    """Periodo coperto da ogni dataset presente. `firma` = tempi di modifica dei file (per la cache)."""
    out = {}
    pg = _p_giornaliero()
    if pg.exists():
        lo, hi = _estremi(pg, "data")
        if lo is not None:
            out["giornaliero"] = (_a_data(lo), _a_data(hi))
    d = C.DATI / "orario"
    if d.is_dir():
        cm = comuni()
        campioni = [_p_orario(i) for i in cm.istat[:3]] if len(cm) else []
        campioni = [p for p in campioni if p.exists()] or sorted(d.glob("*.parquet"))[:3]
        lo = hi = None
        for p in campioni:
            a, b = _estremi(p, "time")
            if a is not None:
                lo = a if lo is None else min(lo, a)
                hi = b if hi is None else max(hi, b)
        if lo is not None:
            out["orario"] = (_a_data(lo), _a_data(hi))
    fe = _file_eurostat()
    if fe:
        periodi = set()
        for p in fe:
            try:
                periodi |= set(pd.read_parquet(p, columns=["time"]).time.astype(str))
            except Exception:
                continue
        date = [_semestre(x) for x in periodi if re.match(r"^\d{4}(-S[12])?$", x)]
        if date:
            out["eurostat"] = (min(a for a, _ in date), max(b for _, b in date))
    return out


def disponibilita() -> dict:
    d = C.DATI / "orario"
    firma = (_mtime(_p_giornaliero()), _mtime(d), _mtime(C.DATI / "comuni.parquet"),
             tuple(_mtime(p) for p in _file_eurostat()), int(time.time() // 3600))   # e comunque ogni ora
    return _disponibilita(firma)


def _periodo_dataset(ds: str) -> tuple[dt.date, dt.date] | None:
    disp = disponibilita()
    if ds == "meteo_orario":
        return disp.get("orario")
    if ds in ("meteo_giornaliero", "gradi_giorno", "legge10"):
        return disp.get("giornaliero")
    if ds == "eurostat_pmi":
        return disp.get("eurostat")
    if ds == "comuni":
        return None
    return None


def _presente(ds: str) -> bool:
    if ds == "comuni":
        return len(comuni()) > 0
    if ds == "dpr412":
        return len(comuni()) > 0 and len(dpr412()) > 0
    return _periodo_dataset(ds) is not None


def _grandezze(ds: str) -> list[Grandezza]:
    return {"meteo_orario": ORARIE, "meteo_giornaliero": GIORNALIERE, "gradi_giorno": GRADI_GIORNO}.get(ds, [])


# ---------------------------------------------------------------------------------------------- limiti d'uso
_storia: dict[tuple[str, str], deque] = defaultdict(deque)
_blocco_storia = threading.Lock()
_in_corso = threading.BoundedSemaphore(max(1, C.STORICI_PARALLELE))


def _ip(req: Request) -> str:
    return req.headers.get("x-forwarded-for", req.client.host if req.client else "?").split(",")[0].strip()


def _limita(req: Request, tipo: str) -> None:
    n = {"stima": C.STORICI_STIME_ORA, "estrazione": C.STORICI_ESTRAZIONI_ORA}[tipo]
    with _blocco_storia:
        q = _storia[(_ip(req), tipo)]
        ora = time.time()
        while q and ora - q[0] > 3600:
            q.popleft()
        if len(q) >= n:
            cosa = "download" if tipo == "estrazione" else "stime"
            raise HTTPException(429, f"Troppe richieste: massimo {n} {cosa} all'ora. Riprova più tardi.")
        q.append(ora)


# ---------------------------------------------------------------------------------------------- richiesta
class Area(BaseModel):
    tipo: Literal["tutti", "regione", "provincia", "comuni"] = "regione"
    regione: str | None = Field(default=None, max_length=2)
    provincia: str | None = Field(default=None, max_length=3)
    comuni: list[str] = Field(default_factory=list, max_length=5000)
    solo_capoluoghi: bool = False


class Richiesta(BaseModel):
    dataset: str = Field(max_length=80)
    variabili: list[str] = Field(default_factory=list, max_length=40)
    dal: dt.date | None = None
    al: dt.date | None = None
    frequenza: str | None = Field(default=None, max_length=20)
    copertura: int = Field(default=C.STORICI_COPERTURA, ge=50, le=100)
    area: Area = Field(default_factory=Area)
    stagioni: int = Field(default=10, ge=3, le=30)       # Legge 10: ultime N stagioni complete
    tipologia: str | None = Field(default=None, max_length=20)   # Legge 10
    tassazione: list[str] = Field(default_factory=list, max_length=5)   # Eurostat


@dataclass
class Piano:
    req: Richiesta
    ds: str
    info: dict
    comuni: pd.DataFrame
    frequenza: str | None
    dal: dt.date | None = None          # periodo letto (esteso ai periodi interi e ritagliato sui dati)
    al: dt.date | None = None
    dal_attesi: dt.date | None = None   # periodo su cui si contano i dati attesi
    al_attesi: dt.date | None = None
    grandezze: list[Grandezza] = field(default_factory=list)
    colonne: list[Uscita] = field(default_factory=list)
    righe: int = 0
    avvisi: list[str] = field(default_factory=list)
    serie: list[str] = field(default_factory=list)


def _err(codice: int, testo: str):
    raise HTTPException(codice, testo)


def _controlla_nomi(req: Richiesta) -> None:
    for x in [req.dataset, req.tipologia, req.frequenza, *req.variabili, *req.tassazione]:
        if _vietato(x):
            _err(403, MSG_VIETATO.format(x=x))


def _scegli_comuni(req: Richiesta) -> tuple[pd.DataFrame, list[str]]:
    cm = comuni()
    a, avvisi = req.area, []
    if a.tipo == "tutti":
        sel = cm
    elif a.tipo == "regione":
        r = (a.regione or "").zfill(2)
        if not re.fullmatch(r"\d{2}", r):
            _err(422, "Scegli la regione")
        sel = cm[cm.regione_cod == r]
        if sel.empty:
            _err(422, "Per questa regione il server non ha ancora i dati meteo dei comuni")
    elif a.tipo == "provincia":
        s = (a.provincia or "").upper()
        if not re.fullmatch(r"[A-Z]{2}", s):
            _err(422, "Scegli la provincia")
        sel = cm[cm.provincia.astype(str).str.upper() == s]
        if sel.empty:
            _err(422, f"Nessun comune disponibile nella provincia {s}")
    else:
        codici = [c.strip().zfill(6) for c in a.comuni if c and c.strip()]
        if not codici:
            _err(422, "Scegli almeno un comune")
        if len(codici) > C.STORICI_MAX_COMUNI:
            _err(413, f"Troppi comuni elencati ({len(codici)}): massimo {C.STORICI_MAX_COMUNI}. "
                      "Scegli una provincia o una regione intera.")
        if any(not re.fullmatch(r"\d{6}", c) for c in codici):
            _err(422, "Codici ISTAT dei comuni non validi (servono 6 cifre)")
        sel = cm[cm.istat.isin(codici)]
        mancano = sorted(set(codici) - set(sel.istat))
        if mancano:
            avvisi.append(f"{len(mancano)} comuni scelti non hanno dati sul server e sono esclusi: "
                          + ", ".join(mancano[:10]) + (" …" if len(mancano) > 10 else ""))
        if sel.empty:
            _err(422, "I comuni scelti non hanno dati sul server")
    if a.solo_capoluoghi:
        sel = sel[sel.capoluogo]
        if sel.empty:
            _err(422, "Nella selezione non ci sono capoluoghi di provincia")
    return sel.reset_index(drop=True), avvisi


def _estendi(dal: dt.date, al: dt.date, freq: str) -> tuple[dt.date, dt.date]:
    """Per le frequenze aggregate il periodo comprende mesi, anni o stagioni interi."""
    if freq == "mensile":
        fine = (pd.Timestamp(al) + pd.offsets.MonthEnd(0)).date()
        return dal.replace(day=1), fine
    if freq == "annuale":
        return dt.date(dal.year, 1, 1), dt.date(al.year, 12, 31)
    if freq == "stagione":
        a = dal.year if dal.month >= 7 else dal.year - 1
        b = al.year if al.month >= 7 else al.year - 1
        return dt.date(a, 7, 1), dt.date(b + 1, 6, 30)
    return dal, al


def _chiave(t: pd.Series, freq: str) -> pd.Series:
    """Etichetta del periodo per date o orari locali (codici numerici, poi solo le etichette distinte:
    strftime su centinaia di migliaia di orari con fuso è troppo lento)."""
    a, m = t.dt.year.to_numpy(), t.dt.month.to_numpy()
    if freq == "giornaliera":
        k, fmt = a * 10000 + m * 100 + t.dt.day.to_numpy(), lambda x: f"{x // 10000}-{x // 100 % 100:02d}-{x % 100:02d}"
    elif freq == "mensile":
        k, fmt = a * 100 + m, lambda x: f"{x // 100}-{x % 100:02d}"
    elif freq == "annuale":
        k, fmt = a, str
    elif freq == "stagione":
        k, fmt = np.where(m >= 7, a, a - 1), lambda x: f"{x}/{(x + 1) % 100:02d}"
    else:
        raise ValueError(freq)
    uniche, inv = np.unique(k, return_inverse=True)
    etichette = np.array([fmt(int(x)) for x in uniche], dtype=object)
    return pd.Series(etichette[inv.reshape(-1)], index=t.index)


def _attesi(p: Piano, orario: bool) -> pd.Series:
    """Numero di dati attesi (giorni od ore) per ogni periodo, sul periodo esteso."""
    if orario:
        t = pd.date_range(pd.Timestamp(p.dal_attesi, tz=ROMA), pd.Timestamp(p.al_attesi + dt.timedelta(days=1), tz=ROMA),
                          freq="h", inclusive="left")
    else:
        t = pd.date_range(p.dal_attesi, p.al_attesi, freq="D")
    return _chiave(pd.Series(t), p.frequenza).value_counts()


def pianifica(req: Richiesta) -> Piano:
    _controlla_nomi(req)
    ds = req.dataset
    if ds not in DATASET:
        _err(404, f"Dataset «{ds}» inesistente")
    if not _presente(ds):
        _err(404, f"Il dataset «{DATASET[ds]['titolo']}» non è ancora disponibile su questo server")
    info = DATASET[ds]
    if info["area"]:
        cm, avvisi = _scegli_comuni(req)
    else:
        cm, avvisi = comuni().head(0), []
    p = Piano(req, ds, info, cm, None, avvisi=avvisi)
    n = len(cm)

    if ds in ("meteo_orario", "meteo_giornaliero", "gradi_giorno"):
        freq = req.frequenza or info["nativa"]
        if freq not in info["frequenze"]:
            _err(422, f"Frequenza «{freq}» non disponibile per questo dataset: "
                      + ", ".join(info["frequenze"]))
        p.frequenza = freq
        tutte = {g.id: g for g in _grandezze(ds)}
        scelte = req.variabili or [g.id for g in tutte.values() if g.predefinita]
        ignote = [v for v in scelte if v not in tutte]
        if ignote:
            _err(422, "Grandezze non disponibili per questo dataset: " + ", ".join(ignote))
        p.grandezze = [tutte[v] for v in tutte if v in scelte]   # nell'ordine del catalogo
        inizio, fine = _periodo_dataset(ds)
        dal, al = req.dal or inizio, req.al or fine
        if dal > al:
            _err(422, "La data di inizio è successiva a quella di fine")
        p.dal_attesi, p.al_attesi = _estendi(dal, al, freq)
        p.dal, p.al = max(p.dal_attesi, inizio), min(p.al_attesi, fine)
        if p.dal > p.al:
            _err(422, f"Nessun dato nel periodo scelto: il dataset va dal {inizio:%d/%m/%Y} al {fine:%d/%m/%Y}")
        giorni = (p.al - p.dal).days + 1
        nativa = freq == info["nativa"]
        if freq == "oraria":
            p.colonne = list(_COL_TEMPO_ORARIA)
        elif freq == "giornaliera":   # dal dato orario il giorno è un'aggregazione: servono i conteggi
            p.colonne = [*_COL_TEMPO_GIORNO, *(_COL_CONTEGGI if not nativa else [])]
        else:
            p.colonne = list(_COL_TEMPO_AGG)
        for g in p.grandezze:
            p.colonne += list(g.base if nativa else g.aggregate)
        if freq == "oraria":
            p.righe = n * giorni * 24
        elif freq == "giornaliera":
            p.righe = n * giorni
        else:
            chiavi = _chiave(pd.Series(pd.date_range(p.dal, p.al, freq="D")), freq)
            p.righe = n * chiavi.nunique()
        if ds == "meteo_orario" and p.dal < INIZIO_ICON:
            p.avvisi.append("Il meteo orario ICON parte dal 14/11/2022: per gli anni precedenti usa il giornaliero.")
    elif ds == "legge10":
        tip = req.tipologia or "mensili"
        if tip not in LEGGE10_TIPOLOGIE:
            _err(422, "Tipologia non valida: " + ", ".join(LEGGE10_TIPOLOGIE))
        if n > C.STORICI_MAX_COMUNI_LEGGE10:
            _err(413, f"Troppi comuni per i dati della Legge 10 ({n}): massimo {C.STORICI_MAX_COMUNI_LEGGE10} per "
                      "file. Scegli una provincia, solo i capoluoghi o un elenco di comuni.")
        inizio, fine = _periodo_dataset(ds)
        p.al = min(req.al or fine, fine)
        p.dal = inizio
        p.frequenza = None
        stagioni_disp = max(0, (p.al - inizio).days // 365)
        k = min(req.stagioni, stagioni_disp) or 1
        p.colonne = list(_COL_LEGGE10[tip])
        p.righe = n * {"mensili": 12, "riepilogo": 1, "stagioni": k}[tip]
    elif ds == "comuni":
        p.colonne = [c for c in _COL_COMUNI if c.sorgente in cm.columns or c.sorgente in ("", "capoluogo")]
        p.righe = n
    elif ds == "dpr412":
        p.colonne = list(_COL_DPR412)
        d = dpr412()
        p.righe = int(d.istat.isin(cm.istat).sum())
    elif ds == "eurostat_pmi":
        file = {f.stem: f for f in _file_eurostat()}
        scelte = req.variabili or list(file)
        ignote = [v for v in scelte if v not in file]
        if ignote:
            _err(422, "Serie Eurostat non disponibili: " + ", ".join(ignote))
        tasse = req.tassazione or list(TASSAZIONE)
        if any(t not in TASSAZIONE for t in tasse):
            _err(422, "Livello di tassazione non valido: " + ", ".join(TASSAZIONE))
        p.req = req.model_copy(update={"tassazione": tasse})
        p.serie = scelte
        inizio, fine = _periodo_dataset(ds)
        p.dal, p.al = max(req.dal or inizio, inizio), min(req.al or fine, fine)
        if p.dal > p.al:
            _err(422, f"Nessun dato nel periodo scelto: i prezzi vanno dal {inizio:%m/%Y} al {fine:%m/%Y}")
        df = _eurostat(p)
        p.colonne = [c for c in _COL_EUROSTAT if c.nome in df.columns]
        p.righe = len(df)
    return p


# colonne fisse
_COL_LUOGO = [Uscita("istat", "", "codice ISTAT del comune", larghezza=7),
              Uscita("comune", "", "nome del comune", larghezza=14),
              Uscita("provincia", "", "sigla della provincia", larghezza=3)]
_COL_TEMPO_ORARIA = [*_COL_LUOGO, Uscita("ora_utc", "UTC", "inizio dell'ora in tempo universale", larghezza=17),
                     Uscita("ora_locale", "ora italiana", "inizio dell'ora in ora civile italiana (CET/CEST)",
                            larghezza=17)]
_COL_TEMPO_GIORNO = [*_COL_LUOGO, Uscita("data", "", "giorno civile italiano", larghezza=11)]
_COL_CONTEGGI = [Uscita("dati_validi", "n", "dati presenti nel periodo (ore o giorni)", larghezza=4),
                 Uscita("dati_attesi", "n", "dati attesi nel periodo (ore o giorni)", larghezza=4)]
_COL_TEMPO_AGG = [*_COL_LUOGO, Uscita("periodo", "", "mese (AAAA-MM), anno (AAAA) o stagione (AAAA/AA, 1 luglio – 30 giugno)",
                                      larghezza=8), *_COL_CONTEGGI]
_COL_COMUNI = [Uscita("istat", "", "codice ISTAT del comune", "istat", larghezza=7),
               Uscita("comune", "", "nome del comune", "nome", larghezza=14),
               Uscita("belfiore", "", "codice catastale (Belfiore)", "belfiore", larghezza=5),
               Uscita("provincia", "", "sigla della provincia", "provincia", larghezza=3),
               Uscita("regione_cod", "", "codice ISTAT della regione", "regione_cod", larghezza=3),
               Uscita("regione", "", "regione", "regione", larghezza=14),
               Uscita("capoluogo", "", "capoluogo di provincia (sì/no)", "capoluogo", larghezza=3),
               Uscita("lat", "°", "latitudine del punto meteo (WGS84)", "lat", dec=5, larghezza=9),
               Uscita("lon", "°", "longitudine del punto meteo (WGS84)", "lon", dec=5, larghezza=9),
               Uscita("metodo_punto", "", "come è stato scelto il punto meteo", "metodo_punto", larghezza=15)]
_COL_LEGGE10 = {
    "mensili": [*_COL_LUOGO, Uscita("stagioni", "", "stagioni di riscaldamento usate", larghezza=16),
                Uscita("mese", "", "mese (1–12)", larghezza=3),
                Uscita("t_media", "°C", "temperatura media mensile", dec=1),
                Uscita("t_min_media", "°C", "media delle temperature minime giornaliere", dec=1),
                Uscita("t_max_media", "°C", "media delle temperature massime giornaliere", dec=1),
                Uscita("h_globale", "MJ/m² giorno", "irradiazione globale giornaliera media su piano orizzontale", dec=1),
                Uscita("h_diffusa", "MJ/m² giorno", "irradiazione diffusa giornaliera media su piano orizzontale", dec=1),
                Uscita("pressione_vapore", "Pa", "pressione parziale del vapore (Magnus)", dec=0),
                Uscita("ur_media", "%", "umidità relativa media", dec=0),
                Uscita("vento", "m/s", "velocità media del vento a 10 m", dec=1)],
    "riepilogo": [*_COL_LUOGO, Uscita("stagioni", "", "stagioni di riscaldamento usate", larghezza=16),
                  Uscita("n_stagioni", "n", "numero di stagioni complete", dec=0, larghezza=3),
                  Uscita("gg_periodo_zona", "°C·giorno", "gradi giorno (base 20 °C) sul periodo di accensione della zona", dec=0),
                  Uscita("zona_riferimento", "", "zona del periodo di accensione usato", larghezza=2),
                  Uscita("zona_risultante", "", "zona climatica che risulta dai GG (DPR 412/93, art. 2)", larghezza=2),
                  Uscita("gg_anno_intero", "°C·giorno", "gradi giorno sull'intera stagione 1 luglio – 30 giugno", dec=0),
                  Uscita("gg_minimo", "°C·giorno", "stagione con meno gradi giorno", dec=0),
                  Uscita("gg_massimo", "°C·giorno", "stagione con più gradi giorno", dec=0),
                  Uscita("tendenza_gg_decennio", "°C·giorno/10 anni", "tendenza dei GG (servono almeno 5 stagioni)", dec=0),
                  Uscita("t_progetto_invernale", "°C", "temperatura esterna di progetto invernale (stima statistica)", dec=1),
                  Uscita("t_progetto_estiva", "°C", "temperatura esterna di progetto estiva (stima statistica)", dec=1),
                  Uscita("metodo_progetto", "", "come sono state stimate le temperature di progetto", larghezza=60)],
    "stagioni": [*_COL_LUOGO, Uscita("stagione", "", "stagione di riscaldamento (1 luglio – 30 giugno)", larghezza=7),
                 Uscita("zona_riferimento", "", "zona del periodo di accensione usato", larghezza=2),
                 Uscita("gg_periodo_zona", "°C·giorno", "gradi giorno (base 20 °C) sul periodo di accensione della zona",
                        dec=0)],
}
_COL_DPR412 = [*_COL_LUOGO,
               Uscita("gradi_giorno", "°C·giorno", "gradi giorno dell'Allegato A del DPR 412/93", dec=0),
               Uscita("zona", "", "zona climatica (DPR 412/93, art. 2)", larghezza=2),
               Uscita("altitudine", "m", "altitudine della casa comunale sul livello del mare", dec=0, larghezza=4),
               Uscita("nome_nel_dpr", "", "nome del comune come scritto nel DPR", larghezza=14),
               Uscita("tipo", "", "diretto, incorporazione, fusione (una riga per comune d'origine) o assente",
                      larghezza=8),
               Uscita("nota", "", "cambi di nome, fusioni, refusi del DPR", larghezza=20)]
_COL_EUROSTAT = [Uscita("serie", "", "serie (file Eurostat)", larghezza=15),
                 Uscita("prodotto", "", "codice del prodotto Eurostat", larghezza=4),
                 Uscita("fascia_codice", "", "fascia di consumo annuo (codice Eurostat)", larghezza=12),
                 Uscita("fascia", "", "fascia di consumo annuo", larghezza=30),
                 Uscita("tassazione", "", "X_TAX senza imposte, X_VAT IVA esclusa, I_TAX tutto incluso", larghezza=5),
                 Uscita("unita", "", "unità di misura Eurostat", larghezza=6),
                 Uscita("valuta", "", "valuta", larghezza=4),
                 Uscita("paese", "", "paese", larghezza=3),
                 Uscita("semestre", "", "semestre (AAAA-S1 gennaio–giugno, AAAA-S2 luglio–dicembre)", larghezza=8),
                 Uscita("valore", "valuta/unità", "prezzo nell'unità di Eurostat", dec=4),
                 Uscita("valore_eur_mwh", "€/MWh", "prezzo in €/MWh (solo se l'unità è il kWh)", dec=2)]


# ---------------------------------------------------------------------------------------------- lettura
def _leggi_giornaliero(codici: list[str], colonne: list[str], dal: dt.date, al: dt.date) -> pd.DataFrame:
    """Solo le colonne e le righe richieste (filtri sul parquet, che è ordinato per comune e data)."""
    p = _p_giornaliero()
    cols = ["istat", "data", *[c for c in dict.fromkeys(colonne) if c not in ("istat", "data")]]
    filtri = [("istat", "in", list(codici))]
    try:
        t = pq.read_table(p, columns=cols, filters=filtri + [("data", ">=", dal), ("data", "<=", al)])
    except (pa.ArrowInvalid, pa.ArrowNotImplementedError, pa.ArrowTypeError, TypeError):
        t = pq.read_table(p, columns=cols, filters=filtri)   # colonna data di tipo diverso da date32
    df = t.to_pandas()
    df["istat"] = df.istat.astype(str)
    d = pd.to_datetime(df["data"])
    if getattr(d.dt, "tz", None) is not None:
        d = d.dt.tz_convert(ROMA).dt.tz_localize(None)
    df["data"] = d.dt.normalize()
    df = df[(df.data >= pd.Timestamp(dal)) & (df.data <= pd.Timestamp(al))]
    return df.sort_values(["istat", "data"]).reset_index(drop=True)


def _leggi_orario(istat: str, colonne: list[str], dal: dt.date, al: dt.date) -> pd.DataFrame:
    p = _p_orario(istat)
    if not p.exists():
        return pd.DataFrame(columns=["time", *colonne])
    presenti = set(pq.read_schema(p).names)
    df = pd.read_parquet(p, columns=["time", *[c for c in colonne if c in presenti]])
    for c in colonne:
        if c not in df:
            df[c] = np.nan
    t = pd.to_datetime(df["time"], utc=True)
    a = pd.Timestamp(dal, tz=ROMA).tz_convert("UTC")
    b = pd.Timestamp(al + dt.timedelta(days=1), tz=ROMA).tz_convert("UTC")
    df["time"] = t
    df = df[(t >= a) & (t < b)].drop_duplicates("time").sort_values("time")
    return df.reset_index(drop=True)


def _a_blocchi(codici: list[str], righe_per_comune: int, massimo: int = 400_000) -> Iterator[list[str]]:
    k = max(1, massimo // max(1, righe_per_comune))
    for i in range(0, len(codici), k):
        yield codici[i:i + k]


def _luogo(df: pd.DataFrame, cm: pd.DataFrame) -> pd.DataFrame:
    nomi = cm.set_index("istat")
    df.insert(0, "provincia", df.istat.map(nomi.provincia).astype(str) if "provincia" in nomi else "")
    df.insert(0, "comune", df.istat.map(nomi.nome) if "nome" in nomi else "")
    istat = df.pop("istat")
    df.insert(0, "istat", istat)
    return df


def _arrotonda(df: pd.DataFrame, colonne: list[Uscita]) -> pd.DataFrame:
    for c in colonne:
        if c.nome in df and c.dec is not None and pd.api.types.is_float_dtype(df[c.nome]):
            df[c.nome] = df[c.nome].astype("float64").round(c.dec)
    return df[[c.nome for c in colonne if c.nome in df]]


def _aggrega(df: pd.DataFrame, t_locale: pd.Series, p: Piano, attesi: pd.Series) -> pd.DataFrame:
    """Aggrega per comune e periodo; i periodi con meno della copertura minima restano vuoti."""
    usc = [u for g in p.grandezze for u in g.aggregate]
    x = df.assign(periodo=_chiave(t_locale, p.frequenza).to_numpy())
    gr = x.groupby(["istat", "periodo"], sort=True)
    spec = {u.nome: (u.sorgente, u.agg) for u in usc}
    out = gr.agg(**spec) if spec else gr.size().to_frame("_n").drop(columns="_n")
    out["dati_validi"] = gr.size()
    out = out.reset_index()
    out["dati_attesi"] = out.periodo.map(attesi).fillna(0).astype(int)
    valido = out.dati_validi >= np.ceil(out.dati_attesi * p.req.copertura / 100.0)
    for u in usc:
        out[u.nome] = out[u.nome].astype("float64") * u.scala
        out.loc[~valido, u.nome] = np.nan
    return out


def _gen_orario(p: Piano) -> Iterator[pd.DataFrame]:
    sorgenti = list(dict.fromkeys(u.sorgente for g in p.grandezze for u in g.base))
    attesi = _attesi(p, orario=True) if p.frequenza != "oraria" else None
    for r in p.comuni.itertuples():
        df = _leggi_orario(r.istat, sorgenti, p.dal, p.al)
        if df.empty:
            continue
        df.insert(0, "istat", r.istat)
        t_loc = df.time.dt.tz_convert(ROMA)
        if p.frequenza == "oraria":
            out = pd.DataFrame({"istat": df.istat, "ora_utc": df.time.dt.strftime("%Y-%m-%dT%H:%MZ"),
                                "ora_locale": t_loc.dt.strftime("%Y-%m-%d %H:%M")})
            for g in p.grandezze:
                for u in g.base:
                    out[u.nome] = df[u.sorgente].astype("float64")
        else:
            out = _aggrega(df, t_loc, p, attesi)
            if p.frequenza == "giornaliera":
                out = out.rename(columns={"periodo": "data"})
        yield _arrotonda(_luogo(out, p.comuni), p.colonne)


def _gen_giornaliero(p: Piano) -> Iterator[pd.DataFrame]:
    sorgenti = list(dict.fromkeys(u.sorgente for g in p.grandezze for u in g.base))
    giorni = (p.al - p.dal).days + 1
    nativa = p.frequenza == "giornaliera"
    attesi = None if nativa else _attesi(p, orario=False)
    for blocco in _a_blocchi(list(p.comuni.istat), giorni):
        df = _leggi_giornaliero(blocco, sorgenti, p.dal, p.al)
        if df.empty:
            continue
        if nativa:
            out = pd.DataFrame({"istat": df.istat, "data": df.data.dt.strftime("%Y-%m-%d")})
            for g in p.grandezze:
                for u in g.base:
                    out[u.nome] = df[u.sorgente].astype("float64")
        else:
            out = _aggrega(df, df.data, p, attesi)
        yield _arrotonda(_luogo(out, p.comuni), p.colonne)


L10_COLONNE = ["t_media", "t_min", "t_max", "ur_media", "rad_globale_kwh_m2", "rad_diffusa_kwh_m2", "vento_medio_kmh"]


def _gen_legge10(p: Piano) -> Iterator[pd.DataFrame]:
    from meteo_storico import legge10
    tip = p.req.tipologia or "mensili"
    scartati = []
    for blocco in _a_blocchi(list(p.comuni.istat), 366 * 31):
        df = _leggi_giornaliero(blocco, L10_COLONNE, p.dal, p.al)
        righe = []
        for istat, g in df.groupby("istat", sort=True):
            try:
                r = legge10.calcola(g.drop(columns="istat").reset_index(drop=True), p.req.stagioni,
                                    fino_a=str(p.al))
            except ValueError:
                scartati.append(istat)
                continue
            st = r["stagioni"]
            periodo = f"{st[0]} – {st[-1]}"
            gg = r["gradi_giorno"]
            if tip == "mensili":
                for m in r["mensili"]:
                    righe.append({"istat": istat, "stagioni": periodo, "mese": int(m["mese"]),
                                  "t_media": m["t_media"], "t_min_media": m["t_min_media"],
                                  "t_max_media": m["t_max_media"], "h_globale": m["h_globale_mj"],
                                  "h_diffusa": m["h_diffusa_mj"], "pressione_vapore": m["pv_pa"],
                                  "ur_media": m["ur_media"], "vento": m["vento_ms"]})
            elif tip == "riepilogo":
                tp = r["temperatura_progetto"]
                righe.append({"istat": istat, "stagioni": periodo, "n_stagioni": len(st),
                              "gg_periodo_zona": gg["gg_periodo_zona"], "zona_riferimento": gg["zona_di_riferimento"],
                              "zona_risultante": gg["zona_risultante"], "gg_anno_intero": gg["gg_anno_intero"],
                              "gg_minimo": gg["minimo"], "gg_massimo": gg["massimo"],
                              "tendenza_gg_decennio": gg["tendenza_gg_per_decennio"],
                              "t_progetto_invernale": tp["invernale"], "t_progetto_estiva": tp["estiva"],
                              "metodo_progetto": tp["metodo"]})
            else:
                for s, v in gg["per_stagione"].items():
                    righe.append({"istat": istat, "stagione": s, "zona_riferimento": gg["zona_di_riferimento"],
                                  "gg_periodo_zona": v})
        if righe:
            out = pd.DataFrame(righe)
            for c in out.columns:
                if c not in ("istat", "stagioni", "stagione", "zona_riferimento", "zona_risultante", "metodo_progetto"):
                    out[c] = pd.to_numeric(out[c], errors="coerce").astype("float64")
            if "mese" in out:
                out["mese"] = out.mese.astype(int)
            if "n_stagioni" in out:
                out["n_stagioni"] = out.n_stagioni.astype(int)
            yield _arrotonda(_luogo(out, p.comuni), p.colonne)
    if scartati:
        p.avvisi.append(f"{len(scartati)} comuni esclusi perché non hanno ancora una stagione completa: "
                        + ", ".join(scartati[:10]) + (" …" if len(scartati) > 10 else ""))


def _gen_comuni(p: Piano) -> Iterator[pd.DataFrame]:
    cm = p.comuni
    out = pd.DataFrame({c.nome: cm[c.sorgente] for c in p.colonne if c.sorgente in cm.columns})
    if "capoluogo" in out:
        out["capoluogo"] = np.where(cm.capoluogo, "sì", "no")
    yield _arrotonda(out, p.colonne)


def _gen_dpr412(p: Piano) -> Iterator[pd.DataFrame]:
    d = dpr412()
    d = d[d.istat.isin(p.comuni.istat)].copy()
    out = _luogo(d, p.comuni)
    yield _arrotonda(out[[c.nome for c in p.colonne]], p.colonne)


def _eurostat(p: Piano) -> pd.DataFrame:
    parti = []
    for f in _file_eurostat():
        if f.stem not in p.serie:
            continue
        d = pd.read_parquet(f)
        d["serie"] = f.stem
        parti.append(d)
    if not parti:
        return pd.DataFrame(columns=[c.nome for c in _COL_EUROSTAT])
    d = pd.concat(parti, ignore_index=True)
    if "tax" in d:
        d = d[d.tax.isin(p.req.tassazione)]
    d["time"] = d.time.astype(str)
    d = d[d.time.str.match(r"^\d{4}(-S[12])?$")]
    dentro = np.array([b >= p.dal and a <= p.al for a, b in map(_semestre, d.time)], dtype=bool)
    d = d.loc[dentro]
    nomi = {"serie": "serie", "product": "prodotto", "nrg_cons": "fascia_codice", "fascia_consumo": "fascia",
            "tax": "tassazione", "unit": "unita", "currency": "valuta", "geo": "paese", "time": "semestre",
            "valore": "valore"}
    out = d[[k for k in nomi if k in d]].rename(columns=nomi)
    out["valore"] = pd.to_numeric(out.valore, errors="coerce").astype("float64")
    if "unita" in out:
        kwh = out.unita.astype(str).str.upper().eq("KWH")
        if kwh.any():
            out["valore_eur_mwh"] = np.where(kwh, out.valore * 1000, np.nan)
    ordine = [c for c in ("serie", "tassazione", "fascia_codice", "semestre") if c in out]
    return out.sort_values(ordine).reset_index(drop=True)


def _gen_eurostat(p: Piano) -> Iterator[pd.DataFrame]:
    yield _arrotonda(_eurostat(p), p.colonne)


def blocchi(p: Piano) -> Iterator[pd.DataFrame]:
    if p.ds == "meteo_orario":
        return _gen_orario(p)
    if p.ds in ("meteo_giornaliero", "gradi_giorno"):
        return _gen_giornaliero(p)
    return {"legge10": _gen_legge10, "comuni": _gen_comuni, "dpr412": _gen_dpr412,
            "eurostat_pmi": _gen_eurostat}[p.ds](p)


# ---------------------------------------------------------------------------------------------- stima e limiti
def _byte_riga(p: Piano) -> int:
    return sum(c.larghezza for c in p.colonne) + len(p.colonne) + 1


def _mb(n: float) -> str:
    if n < 1024 * 1024:
        return f"{max(1, round(n / 1024)):,} kB".replace(",", ".")
    return f"{n / 1024 / 1024:,.1f} MB".replace(",", "X").replace(".", ",").replace("X", ".")


def _suggerimenti(p: Piano, limite: int, formato: str) -> list[str]:
    s = []
    n = max(1, len(p.comuni))
    if p.frequenza == "oraria":
        s.append("scegli una frequenza giornaliera o mensile invece di oraria")
    elif p.frequenza == "giornaliera":
        s.append("scegli una frequenza mensile o annuale invece di giornaliera")
    if p.frequenza and p.dal and p.al:
        per_giorno = {"oraria": 24, "giornaliera": 1}.get(p.frequenza)
        if per_giorno:
            giorni = limite // (n * per_giorno)
            if giorni >= 1:
                s.append(f"accorcia il periodo a circa {giorni:,} giorni".replace(",", ".") + f" con {n} comuni")
    if n > 1:
        s.append(f"riduci i comuni (ora {n}): una provincia, solo i capoluoghi o un elenco")
    if formato == "xlsx":
        s.append(f"scegli il CSV, che arriva fino a {C.STORICI_MAX_RIGHE_CSV:,} righe".replace(",", "."))
    s.append("dividi la richiesta in più file, per esempio un anno o una provincia per volta")
    return s


def valuta(p: Piano, formato: str) -> dict:
    righe = int(p.righe)
    colonne = len(p.colonne)
    byte_csv = righe * _byte_riga(p)
    stima = byte_csv if formato == "csv" else int(byte_csv * 0.75) + 8000   # misurato: Excel ≈ 3/4 del CSV
    limite = C.STORICI_MAX_RIGHE_CSV if formato == "csv" else C.STORICI_MAX_RIGHE_XLSX
    motivo = None
    if righe > limite:
        motivo = (f"La richiesta produrrebbe circa {righe:,} righe: il limite per un file "
                  f"{'CSV' if formato == 'csv' else 'Excel'} è {limite:,}.").replace(",", ".")
    elif formato == "xlsx" and righe * colonne > C.STORICI_MAX_CELLE_XLSX:
        motivo = (f"Il file Excel avrebbe circa {righe * colonne:,} celle: il limite è "
                  f"{C.STORICI_MAX_CELLE_XLSX:,}. Togli qualche grandezza o scegli il CSV.").replace(",", ".")
        limite = C.STORICI_MAX_CELLE_XLSX // max(1, colonne)
    elif righe == 0:
        motivo = "Nessuna riga da estrarre con questi filtri."
    return {
        "dataset": p.ds, "titolo": p.info["titolo"], "formato": formato,
        "righe": righe, "colonne": colonne, "byte": stima, "dimensione": _mb(stima),
        "comuni": len(p.comuni) if p.info["area"] else None,
        "dal": p.dal.isoformat() if p.dal else None, "al": p.al.isoformat() if p.al else None,
        "frequenza": p.frequenza, "limite_righe": limite,
        "consentito": motivo is None, "motivo": motivo,
        "suggerimenti": _suggerimenti(p, limite, formato) if motivo and righe else [],
        "avvisi": p.avvisi, "intestazioni": [_intestazione(c) for c in p.colonne],
        "fonte": p.info["fonte"], "licenza": p.info["licenza"][0], "licenza_url": p.info["licenza"][1],
    }


# ---------------------------------------------------------------------------------------------- scrittura
def _intestazione(c: Uscita) -> str:
    return f"{c.nome} ({c.unita})" if c.unita else c.nome


def _descr_area(p: Piano) -> str:
    a = p.req.area
    if not p.info["area"]:
        return "Italia"
    cm = p.comuni
    testo = {"tutti": "tutti i comuni disponibili",
             "regione": f"regione {cm.regione.iloc[0] if len(cm) and 'regione' in cm else a.regione}",
             "provincia": f"provincia di {a.provincia}",
             "comuni": "comuni scelti: " + ", ".join(cm.nome.head(30)) + (" …" if len(cm) > 30 else "")}[a.tipo]
    if a.solo_capoluoghi:
        testo += " (solo capoluoghi di provincia)"
    return f"{testo} – {len(cm)} comuni"


def _filtri(p: Piano) -> list[tuple[str, str]]:
    r = p.req
    out = [("Area", _descr_area(p))]
    if p.ds in ("meteo_orario", "meteo_giornaliero", "gradi_giorno"):
        out += [("Periodo", f"dal {p.dal:%d/%m/%Y} al {p.al:%d/%m/%Y}"),
                ("Frequenza", FREQUENZE[p.frequenza]),
                ("Grandezze", "; ".join(g.nome for g in p.grandezze))]
        if p.frequenza != p.info["nativa"]:
            out.append(("Copertura minima", f"{r.copertura}% dei dati attesi nel periodo: sotto questa soglia i "
                                            "valori restano vuoti (colonne dati_validi e dati_attesi)"))
            out.append(("Aggregazioni", "medie per temperature, umidità, vento e nuvolosità; somme per radiazione, "
                                        "precipitazione, neve e gradi giorno; minimo e massimo per t_min e t_max"))
    elif p.ds == "legge10":
        out += [("Tipologia", LEGGE10_TIPOLOGIE[r.tipologia or "mensili"]),
                ("Stagioni", f"ultime {r.stagioni} stagioni complete fino al {p.al:%d/%m/%Y}")]
    elif p.ds == "eurostat_pmi":
        out = [("Serie", "; ".join(SERIE_EUROSTAT.get(s, s) for s in p.serie)),
               ("Tassazione", "; ".join(f"{t} {TASSAZIONE[t]}" for t in r.tassazione)),
               ("Periodo", f"dal {p.dal:%m/%Y} al {p.al:%m/%Y}")]
    return out


def _nome_file(p: Piano, formato: str) -> str:
    parti = ["eugenio", p.ds]
    if p.frequenza:
        parti.append(p.frequenza)
    if p.ds == "legge10":
        parti.append(p.req.tipologia or "mensili")
    if p.dal and p.al and p.ds != "legge10":
        parti.append(f"{p.dal:%Y%m%d}-{p.al:%Y%m%d}")
    return "_".join(parti) + "." + formato


def _ora_estrazione() -> str:
    return dt.datetime.now(ROMA).strftime("%d/%m/%Y %H:%M")


def _note(p: Piano) -> list[tuple[str, str]]:
    note = [("Dataset", p.info["titolo"]), *_filtri(p),
            ("Fonte", p.info["fonte"]), ("Licenza", f"{p.info['licenza'][0]} – {p.info['licenza'][1]}"),
            ("Licenza EuGEnio", RIUSO),
            ("Come citare", f"{CITAZIONE} (dati estratti il {dt.datetime.now(ROMA):%d/%m/%Y})"),
            ("Estratto il", f"{_ora_estrazione()} da {C.SITO_URL}/dati.html")]
    if p.info.get("avvertenza"):
        note.append(("Avvertenza", p.info["avvertenza"]))
    return note


def csv_streaming(p: Piano) -> Iterator[bytes]:
    yield ("﻿" + ";".join(_intestazione(c) for c in p.colonne) + "\r\n").encode("utf-8")
    for df in blocchi(p):
        if len(df):
            yield df.to_csv(sep=";", decimal=",", header=False, index=False, lineterminator="\r\n").encode("utf-8")
    righe = ["", *[f"# {k}: {v}" for k, v in _note(p)], *[f"# Nota: {a}" for a in p.avvisi]]
    yield ("\r\n".join(r.replace("\r", " ").replace("\n", " ") for r in righe) + "\r\n").encode("utf-8")


def _valori_xlsx(df: pd.DataFrame) -> Iterator[list]:
    """Righe pronte per openpyxl: date e orari come veri valori di Excel, celle vuote al posto di NaN."""
    cols = []
    for c in df.columns:
        s = df[c]
        if c == "data":
            cols.append([None if pd.isna(x) else x.date() for x in pd.to_datetime(s)])
        elif c in ("ora_utc", "ora_locale"):
            cols.append([None if pd.isna(x) else x.to_pydatetime() for x in pd.to_datetime(s.str.rstrip("Z"))])
        else:
            cols.append(s.astype(object).where(s.notna(), None).tolist())
    yield from (list(r) for r in zip(*cols))


def xlsx_file(p: Piano) -> str:
    """Scrive il file Excel (write-only, poca memoria) in un file temporaneo e ne restituisce il percorso."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Dati")
    ws.freeze_panes = "A2"
    grassetto = Font(bold=True)

    def cella(foglio, v, bold=False, fmt=None):
        c = WriteOnlyCell(foglio, value=v)
        if bold:
            c.font = grassetto
        if fmt:
            c.number_format = fmt
        return c

    ws.append([cella(ws, _intestazione(c), True) for c in p.colonne])
    formati = {"data": "dd/mm/yyyy", "ora_utc": "dd/mm/yyyy hh:mm", "ora_locale": "dd/mm/yyyy hh:mm"}
    for df in blocchi(p):
        nomi = list(df.columns)
        speciali = [i for i, n in enumerate(nomi) if n in formati]
        for r in _valori_xlsx(df):
            for i in speciali:
                if r[i] is not None:
                    r[i] = cella(ws, r[i], fmt=formati[nomi[i]])
            ws.append(r)

    info = wb.create_sheet("Info")
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 110
    info.append([cella(info, "EuGEnio – dati storici", True)])
    info.append([])
    a_capo = Alignment(wrap_text=True, vertical="top")
    for k, v in _note(p):
        c = WriteOnlyCell(info, value=v)
        c.alignment = a_capo
        info.append([cella(info, k, True), c])
    for a in p.avvisi:
        info.append([cella(info, "Nota", True), a])
    info.append([])
    info.append([cella(info, "Colonna", True), cella(info, "Descrizione", True), cella(info, "Unità", True),
                 cella(info, "Aggregazione", True)])
    aggr = {"mean": "media", "sum": "somma", "min": "minimo", "max": "massimo"}
    for c in p.colonne:
        info.append([c.nome, c.descrizione, c.unita, aggr.get(c.agg, "")])
    fd, percorso = tempfile.mkstemp(prefix="eugenio_", suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(percorso)
    except Exception:
        os.unlink(percorso)
        raise
    return percorso


def _leggi_e_cancella(percorso: str) -> Iterator[bytes]:
    try:
        with open(percorso, "rb") as f:
            while b := f.read(1 << 20):
                yield b
    finally:
        try:
            os.unlink(percorso)
        except OSError:
            pass


# ---------------------------------------------------------------------------------------------- rotte
@router.get("/catalogo")
def catalogo():
    """Dataset presenti sul server, con grandezze, frequenze, periodo, fonte e licenza."""
    out = []
    for ds, info in DATASET.items():
        if not _presente(ds):
            continue
        per = _periodo_dataset(ds)
        voce = {"id": ds, "titolo": info["titolo"], "descrizione": info["descrizione"], "fonte": info["fonte"],
                "licenza": info["licenza"][0], "licenza_url": info["licenza"][1], "periodo": info["periodo"],
                "inizio": per[0].isoformat() if per else None, "fine": per[1].isoformat() if per else None,
                "area": info["area"],
                "frequenze": [{"id": f, "nome": FREQUENZE[f]} for f in info["frequenze"]],
                "frequenza_nativa": info.get("nativa"),
                "variabili": [g.catalogo(info.get("nativa")) for g in _grandezze(ds)]}
        if info.get("avvertenza"):
            voce["avvertenza"] = info["avvertenza"]
        if ds == "legge10":
            voce["tipologie"] = [{"id": k, "nome": v} for k, v in LEGGE10_TIPOLOGIE.items()]
            voce["stagioni"] = {"min": 3, "max": 30, "predefinite": 10}
        if ds == "eurostat_pmi":
            voce["variabili"] = [{"id": f.stem, "nome": SERIE_EUROSTAT.get(f.stem, f.stem.replace("_", " ")),
                                  "gruppo": "Serie", "predefinita": True} for f in _file_eurostat()]
            voce["tassazione"] = [{"id": k, "nome": v, "predefinita": k == "X_VAT"} for k, v in TASSAZIONE.items()]
        if ds == "comuni":
            voce["colonne"] = [_intestazione(c) for c in _COL_COMUNI]
        if ds == "dpr412":
            voce["colonne"] = [_intestazione(c) for c in _COL_DPR412]
            voce["fonte_url"] = DPR412_URL
        out.append(voce)
    cm = comuni()
    return {"dataset": out,
            "regioni": sorted(cm.regione_cod.unique().tolist()) if len(cm) else [],
            "comuni": int(len(cm)),
            "copertura_predefinita": C.STORICI_COPERTURA,
            "limiti": {"righe_csv": C.STORICI_MAX_RIGHE_CSV, "righe_xlsx": C.STORICI_MAX_RIGHE_XLSX,
                       "celle_xlsx": C.STORICI_MAX_CELLE_XLSX, "comuni_elenco": C.STORICI_MAX_COMUNI,
                       "comuni_legge10": C.STORICI_MAX_COMUNI_LEGGE10, "download_ora": C.STORICI_ESTRAZIONI_ORA},
            "non_scaricabili": NON_SCARICABILI}


@router.post("/stima")
def stima(request: Request, richiesta: Richiesta, formato: Literal["csv", "xlsx"] = Query("csv")):
    _limita(request, "stima")
    return valuta(pianifica(richiesta), formato)


@router.post("/estrai")
async def estrai(request: Request, richiesta: Richiesta, formato: Literal["csv", "xlsx"] = Query("csv")):
    _controlla_nomi(richiesta)
    p = await run_in_threadpool(pianifica, richiesta)
    v = valuta(p, formato)
    if not v["consentito"]:
        testo = v["motivo"] + (" Per ridurla: " + "; ".join(v["suggerimenti"]) + "." if v["suggerimenti"] else "")
        raise HTTPException(413 if v["righe"] else 422, testo)
    _limita(request, "estrazione")
    if not _in_corso.acquire(blocking=False):
        raise HTTPException(503, "Il server sta già preparando altri file: riprova tra qualche secondo.")
    nome = _nome_file(p, formato)
    intest = {"Content-Disposition": f'attachment; filename="{nome}"', "Cache-Control": "no-store",
              "X-Righe-Stimate": str(v["righe"])}
    if formato == "csv":
        def flusso():
            try:
                yield from csv_streaming(p)
            finally:
                _in_corso.release()
        return StreamingResponse(flusso(), media_type="text/csv; charset=utf-8", headers=intest)
    try:
        percorso = await run_in_threadpool(xlsx_file, p)
    finally:
        _in_corso.release()
    intest["Content-Length"] = str(os.path.getsize(percorso))
    return StreamingResponse(_leggi_e_cancella(percorso),
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers=intest)
