"""Prove dei report scaricabili (servizio/report.py, servizio/report_servizi.py e rotte /api/report).

Si eseguono dalla cartella fase4-server con `python -m pytest tests`. Non servono TiRex-2 né i dati
meteo: per importare app.py i moduli della fase 3 (meteo_storico) sono sostituiti da segnaposto.
"""
from __future__ import annotations

import base64
import io
import struct
import sys
import types
import zlib
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from servizio import report as R                     # noqa: E402
from servizio import report_servizi as RS            # noqa: E402

# ---------------------------------------------------------------------------------------------- strumenti


def png_minimo(larghezza=4, altezza=2) -> bytes:
    """PNG RGB valido costruito a mano (nessun file su disco)."""
    def blocco(tipo, dati):
        return struct.pack(">I", len(dati)) + tipo + dati + struct.pack(">I", zlib.crc32(tipo + dati) & 0xFFFFFFFF)
    righe = b"".join(b"\x00" + b"\xff\x80\x00" * larghezza for _ in range(altezza))
    return (b"\x89PNG\r\n\x1a\n" + blocco(b"IHDR", struct.pack(">IIBBBBB", larghezza, altezza, 8, 2, 0, 0, 0))
            + blocco(b"IDAT", zlib.compress(righe)) + blocco(b"IEND", b""))


SVG = (b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="80" viewBox="0 0 200 80">'
       b'<rect x="10" y="10" width="80" height="60" fill="#3a6ea5"/></svg>')


def leggi_docx(b: bytes) -> str:
    from docx import Document
    d = Document(io.BytesIO(b))
    parti = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            parti += [c.text for c in r.cells]
    parti += [p.text for p in d.sections[0].header.paragraphs]
    parti += [p.text for p in d.sections[0].footer.paragraphs]
    return "\n".join(parti)


def leggi_pdf(b: bytes) -> str:
    from pypdf import PdfReader
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(b)).pages)


def libro_xlsx(b: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(b))


def odf_parti(b: bytes) -> dict[str, str]:
    """Controlli del pacchetto OpenDocument (ODF 1.2): 'mimetype' primo e non compresso, manifest presente;
    restituisce i file XML principali come testo."""
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(b))
    primo = z.infolist()[0]
    assert primo.filename == "mimetype" and primo.compress_type == zipfile.ZIP_STORED
    assert "META-INF/manifest.xml" in z.namelist()
    return {n: z.read(n).decode("utf-8") for n in ("mimetype", "content.xml", "styles.xml", "meta.xml",
                                                    "META-INF/manifest.xml")}


def solo_testo_odf(xml: str) -> str:
    """Testo visibile di un content.xml/styles.xml (senza tag, con le entità risolte)."""
    import html
    import re
    return html.unescape(re.sub(r"<[^>]+>", " ", xml))


def tutto_xlsx(wb) -> str:
    return "\n".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value is not None)


ESEMPIO = {
    "titolo": "Relazione di prova – Forlì",
    "sottotitolo": "Città con accenti: è, à, ù; simboli °C, €, ≥",
    "data": "2026-10-27",
    "sezioni": [
        {"titolo": "Sezione con perché", "blocchi": [
            {"tipo": "paragrafo", "testo": "La temperatura è ≥ 20 °C e il costo è di 12 €."},
            {"tipo": "elenco", "voci": ["prima voce già fatta", "seconda voce: però"]},
            {"tipo": "tabella", "titolo": "Consumi mensili", "intestazioni": ["Giorno", "Energia (kWh)", "Nota"],
             "righe": [["2026-01-01", 1234.5, "capodanno"], ["2026-01-02T08:00:00+01:00", 987, None],
                       ["2026-01-03", 0.25, "=1+1"]], "decimali": [None, 1, None], "nota": "Valori sintetici."},
            {"tipo": "chiave_valore", "titolo": "Parametri", "voci": [
                {"chiave": "Temperatura di progetto", "valore": -5.25, "unita": "°C", "decimali": 1},
                {"chiave": "Zona climatica", "valore": "E"}, {"chiave": "Verificato", "valore": True}]},
            {"tipo": "immagine", "formato": "png", "dati": base64.b64encode(png_minimo()).decode(),
             "didascalia": "Grafico PNG", "larghezza_cm": 4},
            {"tipo": "immagine", "formato": "svg", "dati": base64.b64encode(SVG).decode(), "didascalia": "Grafico SVG"},
            {"tipo": "avvertenza", "testo": "Non sostituisce la UNI 10349."},
        ]},
    ],
    "fonti": [{"testo": "Weather data by Open-Meteo.com (CC BY 4.0)", "url": "https://open-meteo.com/"},
              {"testo": "ISTAT (CC BY 4.0)"}],
}


@pytest.fixture
def doc():
    return R.Documento.model_validate(ESEMPIO)


# ---------------------------------------------------------------------------------------------- formati


def test_numeri_e_date_italiani():
    assert R.numero(1234.5) == "1.234,50"
    assert R.numero(1234567) == "1.234.567"
    assert R.numero(-5.25, 1) == "-5,3"
    assert R.numero(-0.04, 1) == "0,0"
    assert R.testo_cella(None) == "—" and R.testo_cella(True) == "sì"
    assert R.testo_cella("2026-10-27") == "27/10/2026"
    assert R.testo_cella("2026-07-01T12:00:00Z") == "01/07/2026 14:00"       # UTC → ora italiana (estate)
    assert R.testo_cella("037006") == "037006"                               # i codici restano testo


def test_docx(doc):
    b = R.genera(doc, "docx")
    t = leggi_docx(b)
    for atteso in ("Relazione di prova – Forlì", "≥ 20 °C", "12 €", "1.234,5", "02/01/2026 08:00", "-5,3 °C",
                   "prima voce già fatta", "UNI 10349", "Fonti e attribuzioni", "https://open-meteo.com/",
                   "Prodotto da EuGEnio – eugenio.statocritico.it", "Grafico PNG"):
        assert atteso in t, atteso


def test_pdf(doc):
    b = R.genera(doc, "pdf")
    assert b.startswith(b"%PDF")
    t = leggi_pdf(b)
    for atteso in ("Relazione di prova", "Forlì", "≥ 20 °C", "12 €", "1.234,5", "27 ottobre 2026",
                   "-5,3 °C", "già fatta", "UNI 10349", "Weather data by Open-Meteo.com",
                   "Prodotto da EuGEnio", "eugenio.statocritico.it"):
        assert atteso in t, atteso


def test_xlsx(doc):
    import datetime as dt
    wb = libro_xlsx(R.genera(doc, "xlsx"))
    assert wb.sheetnames == ["Info", "Consumi mensili"]
    ws = wb["Consumi mensili"]
    assert [c.value for c in ws[1]] == ["Giorno", "Energia (kWh)", "Nota"]
    assert ws["B2"].value == 1234.5 and ws["B2"].number_format == "#,##0.0"     # numero, non testo
    assert ws["A2"].value == dt.datetime(2026, 1, 1) and ws["A2"].number_format == "dd/mm/yyyy"
    assert ws["A3"].value == dt.datetime(2026, 1, 2, 8, 0)
    assert ws["C4"].value == "=1+1" and ws["C4"].data_type == "s"               # mai formule dai dati
    t = tutto_xlsx(wb)
    for atteso in ("Relazione di prova – Forlì", "Weather data by Open-Meteo.com (CC BY 4.0)", "UNI 10349",
                   "Prodotto da EuGEnio – eugenio.statocritico.it", "Zona climatica"):
        assert atteso in t, atteso


def test_odt(doc):
    p = odf_parti(R.genera(doc, "odt"))
    assert p["mimetype"] == "application/vnd.oasis.opendocument.text"
    t = " ".join(solo_testo_odf(p["content.xml"]).split())
    for atteso in ("Relazione di prova – Forlì", "≥ 20 °C", "12 €", "1.234,5", "02/01/2026 08:00", "-5,3 °C",
                   "prima voce già fatta", "UNI 10349", "Fonti e attribuzioni", "https://open-meteo.com/",
                   "Grafico PNG", R.TITOLO_AVVERTENZE, R.CITAZIONE):
        assert atteso in t, atteso
    assert 'xlink:href="https://open-meteo.com/"' in p["content.xml"]
    stili = p["styles.xml"]
    testa = " ".join(solo_testo_odf(stili).split())
    assert doc.titolo in testa and "Prodotto da EuGEnio – eugenio.statocritico.it" in testa
    assert "<text:page-number" in stili and "<text:page-count" in stili
    assert R.PIE_AVVERTENZA.split(":")[0] in testa
    assert 'fo:page-width="21cm"' in stili and 'fo:page-height="29.7cm"' in stili     # A4 verticale
    assert "#" + R.COLORE_LINK.lower() in (p["content.xml"] + stili).lower()             # collegamenti blu scuro
    assert "Pictures/logo_eugenio.png" in stili and "Pictures/logo_eugenio.png" in p["META-INF/manifest.xml"]
    assert "EuGEnio" in p["meta.xml"]


def test_ods(doc):
    p = odf_parti(R.genera(doc, "ods"))
    assert p["mimetype"] == "application/vnd.oasis.opendocument.spreadsheet"
    x = p["content.xml"]
    assert 'table:name="Info"' in x and 'table:name="Consumi mensili"' in x
    assert 'office:value-type="float" office:value="1234.5"' in x               # numero, non testo
    assert 'office:value-type="date" office:date-value="2026-01-01"' in x
    assert 'office:date-value="2026-01-02T08:00"' in x
    assert "table:formula" not in x and "=1+1" in x                            # mai formule dai dati
    t = " ".join(solo_testo_odf(x).split())
    for atteso in ("Relazione di prova – Forlì", "Weather data by Open-Meteo.com (CC BY 4.0)", "UNI 10349",
                   R.INTESTAZIONE, "Zona climatica", R.TITOLO_AVVERTENZE, R.CITAZIONE):
        assert atteso in t, atteso
    stili = " ".join(solo_testo_odf(p["styles.xml"]).split())
    assert "EuGEnio" in stili and doc.titolo in stili and R.PIE_AVVERTENZA_XLSX in stili
    assert "<text:page-number" in p["styles.xml"]
    import pandas as pd                         # rilettura con il lettore ODF di pandas (odfpy)
    fogli = pd.read_excel(io.BytesIO(R.genera(doc, "ods")), engine="odf", sheet_name=None)
    assert list(fogli) == ["Info", "Consumi mensili"]
    assert fogli["Consumi mensili"]["Energia (kWh)"].tolist()[:2] == [1234.5, 987]


def test_odf_senza_logo(doc, monkeypatch):
    monkeypatch.setattr(R, "LOGO_PNG", None)
    assert "Pictures/logo" not in odf_parti(R.genera(doc, "odt"))["styles.xml"]
    assert "Pictures/logo" not in odf_parti(R.genera(doc, "ods"))["content.xml"]


def test_intestazione_e_avvertenze_docx(doc):
    from docx import Document
    d = Document(io.BytesIO(R.genera(doc, "docx")))
    sez = d.sections[0]
    testata = "".join(p.text for p in sez.header.paragraphs)
    assert doc.titolo in testata and "Prodotto da" not in testata          # in testa: logo e titolo
    assert "w:bottom" in sez.header._element.xml                           # linea sotto l'intestazione
    pie = "\n".join(p.text for p in sez.footer.paragraphs)
    assert pie.startswith("Prodotto da EuGEnio – eugenio.statocritico.it. ")   # poi l'avvertenza, senza a capo
    assert R.PIE_AVVERTENZA in pie and "Pagina" in pie
    assert any(r.target_ref == R.SITO for r in sez.footer.part.rels.values())
    xml = sez.footer._element.xml
    assert " PAGE " in xml and " NUMPAGES " in xml
    corpo = [p.text for p in d.paragraphs]
    assert corpo.count(R.TITOLO_AVVERTENZE) == 1                      # testo completo una volta, in fondo
    assert corpo.index(R.TITOLO_AVVERTENZE) > corpo.index("Fonti e attribuzioni")
    for testo in R.AVVERTENZE:
        assert testo in corpo
    for parola in ("gratuito", "intelligenza artificiale", "TiRex-2", "contatti.html"):
        assert parola in R.PIE_AVVERTENZA and parola in " ".join(R.AVVERTENZE)
    for parola in ("CC BY 4.0", R.CITAZIONE):
        assert parola in R.PIE_AVVERTENZA and parola in R.PIE_AVVERTENZA_XLSX and parola in R.TESTO_LICENZA
    assert R.TESTO_LICENZA in corpo


def test_licenza_e_citazione_testi():
    assert R.CITAZIONE == "Fonte: EuGEnio – eugenio.statocritico.it"
    assert R.TESTO_LICENZA in R.AVVERTENZE
    for parola in ("Creative Commons Attribuzione 4.0", "anche modificati", "citare EuGEnio come fonte",
                   "fonti originali", "Zornade", "atti", "AGPL-3.0", "licenza.html"):
        assert parola in R.TESTO_LICENZA, parola


def test_intestazione_e_avvertenze_pdf(doc):
    from pypdf import PdfReader
    lungo = doc.model_copy(deep=True)
    lungo.sezioni[0].blocchi[2].righe = [["2026-01-01", float(i), "x"] for i in range(150)]   # più pagine
    pagine = [p.extract_text() for p in PdfReader(io.BytesIO(R.genera(lungo, "pdf"))).pages]
    assert len(pagine) >= 3
    for i, testo in enumerate(pagine, 1):
        piatto = " ".join(testo.split())
        assert "Prodotto da EuGEnio" in piatto
        assert "Servizio gratuito e senza garanzia" in piatto and "intelligenza artificiale" in piatto
        assert "contatti.html" in piatto
        assert f"Pagina {i} di {len(pagine)}" in piatto
        assert "CC BY 4.0" in piatto and R.CITAZIONE in piatto           # licenza nel piè di ogni pagina
    tutto = " ".join(" ".join(pagine).split())
    assert tutto.count(R.TITOLO_AVVERTENZE) == 1
    assert "Se trovi degli errori, segnalali" in " ".join(pagine[-1].split())   # avvertenze in fondo
    assert "citare EuGEnio come fonte" in tutto and "report del giorno/mese/anno" in tutto


def test_intestazione_e_avvertenze_xlsx(doc, monkeypatch):
    monkeypatch.setattr(R, "LOGO_PNG", None)   # con il logo l'intestazione scende alla riga 2
    wb = libro_xlsx(R.genera(doc, "xlsx"))
    assert wb["Info"]["A1"].value == R.INTESTAZIONE and wb["Info"]["A1"].hyperlink.target == R.SITO
    for ws in wb.worksheets:
        assert ws.oddHeader.left.text == "EuGEnio" and ws.oddHeader.center.text == doc.titolo
        assert ws.oddFooter.left.text.startswith("Prodotto da EuGEnio")
        assert ws.oddFooter.left.text == R.PIE_AVVERTENZA_XLSX
        assert ws.oddFooter.right.text == "Pagina &P di &N"
        assert len(str(ws.oddFooter)) <= 255 and len(str(ws.oddHeader)) <= 255    # limite di Excel
    info = [c.value for r in wb["Info"].iter_rows() for c in r if c.value is not None]
    assert info.count(R.TITOLO_AVVERTENZE) == 1
    for testo in R.AVVERTENZE:
        assert testo in info


def test_logo_facoltativo(doc, monkeypatch):
    monkeypatch.setattr(R, "LOGO_PNG", png_minimo(30, 10))
    from docx import Document
    d = Document(io.BytesIO(R.genera(doc, "docx")))
    assert "graphic" in d.sections[0].header._element.xml                # immagine nell'intestazione
    assert R.genera(doc, "pdf").startswith(b"%PDF")
    assert libro_xlsx(R.genera(doc, "xlsx"))["Info"].cell(row=2, column=1).value == R.INTESTAZIONE


def test_nome_file(doc):
    assert R.nome_file(doc, "pdf") == "relazione_di_prova_forli_20261027.pdf"
    doc.nome_file = "clima_legge10_037006"
    assert R.nome_file(doc, "xlsx") == "clima_legge10_037006.xlsx"
    doc.nome_file = '../../etc/"passwd"'
    assert R.nome_file(doc, "docx") == "etc_passwd.docx"


def test_pdf_senza_font_dejavu(doc, monkeypatch, tmp_path):
    monkeypatch.setattr(R.C, "REPORT_FONT_DIR", tmp_path)          # cartella vuota: si ripiega su Helvetica
    t = leggi_pdf(R.genera(doc, "pdf"))
    assert "Forlì" in t and ">= 20 °C" in t and "12 EUR" in t


# ---------------------------------------------------------------------------------------------- limiti


def test_limiti(monkeypatch):
    from pydantic import ValidationError
    troppe = dict(ESEMPIO, sezioni=[{"blocchi": [{"tipo": "tabella", "intestazioni": ["a"],
                                                  "righe": [[1]] * (R.C.REPORT_MAX_RIGHE + 1)}]}])
    with pytest.raises(ValidationError):
        R.Documento.model_validate(troppe)
    with pytest.raises(ValidationError):
        R.Documento.model_validate(dict(ESEMPIO, sezioni=[{}] * (R.C.REPORT_MAX_SEZIONI + 1)))
    with pytest.raises(ValidationError):
        R.Documento.model_validate(dict(ESEMPIO, sezioni=[{"blocchi": [{"tipo": "paragrafo",
                                                                        "testo": "x" * (R.C.REPORT_MAX_TESTO + 1)}]}]))
    with pytest.raises(ValidationError):    # riga più lunga delle intestazioni
        R.Documento.model_validate(dict(ESEMPIO, sezioni=[{"blocchi": [{"tipo": "tabella", "intestazioni": ["a"],
                                                                        "righe": [[1, 2]]}]}]))
    cattivo = SVG.replace(b"</svg>", b'<image href="http://esempio.it/x.png"/></svg>')
    with pytest.raises(ValidationError):
        R.Documento.model_validate(dict(ESEMPIO, sezioni=[{"blocchi": [
            {"tipo": "immagine", "formato": "svg", "dati": base64.b64encode(cattivo).decode()}]}]))
    with pytest.raises(ValidationError):    # PNG dichiarato ma dati diversi
        R.Documento.model_validate(dict(ESEMPIO, sezioni=[{"blocchi": [
            {"tipo": "immagine", "formato": "png", "dati": base64.b64encode(b"ciao").decode()}]}]))


# ---------------------------------------------------------------------------------------------- servizi

PREVISIONE = {
    "modalita": "breve", "passo": "ora", "istat": "037006",
    "covariate": ["temperatura", "radiazione", "festivo"],
    "storico": {"t": ["2026-10-26T22:00:00+01:00", "2026-10-26T23:00:00+01:00"], "v": [10.5, None]},
    "previsione": {"t": ["2026-10-27T00:00:00+01:00", "2026-10-27T01:00:00+01:00", "2026-10-27T02:00:00+01:00"],
                   **{q: [9.0 + i / 10, 9.5 + i / 10, 10.0 + i / 10] for i, q in
                      enumerate(["0.1", "0.2", "0.3", "0.4", "0.5", "0.6", "0.7", "0.8", "0.9"])},
                   "fonte_meteo": ["previsione ICON-2I", "previsione ICON-2I", "miscela ICON-2I/IFS"]},
    "avvisi": ["Oltre le 72 ore il meteo proviene dal modello globale ECMWF IFS (0,25°)."],
    "verifica": {"reale": [9.8, 10.1, None], "mae": 0.31, "wape": 0.042, "bias": -0.1, "copertura_80": 0.875},
    "richiesta": {"comune": "Bologna", "modalita": "breve", "orizzonte": 3, "covariate": "temperatura,radiazione",
                  "tipo": "energia", "unita": "kWh"},
}

LEGGE10 = {
    "istat": "037006", "comune": "Bologna", "stagioni": ["2016/17", "2017/18", "2024/25"], "giorni": 1653,
    "gradi_giorno": {"gg_periodo_zona": 2071.4, "zona_di_riferimento": "E", "zona_risultante": "E",
                     "gg_anno_intero": 2150.2, "minimo": 1890.0, "massimo": 2290.3, "tendenza_gg_per_decennio": -84.2,
                     "per_stagione": {"2016/17": 2290.3, "2017/18": 2101.0, "2024/25": 1890.0},
                     "gg_per_periodo_di_ogni_zona": {"D": 1720.5, "E": 2071.4, "F": 2400.1}},
    "temperatura_progetto": {"invernale": -3.84, "estiva": 33.2, "metodo": "percentile 99,6 % delle medie orarie"},
    "mensili": [{"mese": m, "t_media": 2.5 + m, "t_min_media": m - 1.0, "t_max_media": m + 6.0, "h_globale_mj": 4.2 + m,
                 "h_diffusa_mj": 2.1, "pv_pa": 700 + m, "ur_media": 80.0, "vento_ms": 2.3} for m in range(1, 13)],
    "avvertenza": "Valori dal clima reale: non sostituiscono UNI 10349 e DPR 412/93.",
    "fonte": "Elaborazione su dati Open-Meteo.com (ERA5/Copernicus corretto), CC BY 4.0",
}

PATRIMONIO = {
    "righe": [
        {"id": "E1", "denominazione": "Scuola Pascoli", "indirizzo": "Via Indipendenza", "civico": "10",
         "comune": "Bologna", "esito": "ok", "indirizzo_trovato": "Via dell'Indipendenza 10, Bologna",
         "fonte_punto": "ANNCSU", "comune_catasto": "A944", "foglio_catasto": 200, "mappale": "45",
         "area_m2": 1234.56, "tipo_particella": "fabbricato", "lat": 44.4985, "lon": 11.3430, "prezzo_pun": 0.12},
        {"id": "E2", "denominazione": "Palestra Forlì", "indirizzo": "Via Inesistente", "comune": "Forlì",
         "esito": "non_trovato", "lat": None, "lon": None, "nota": "indirizzo non trovato"},
        {"id": "E3", "esito": "verificare", "lat": 44.2, "lon": 12.04},
    ],
    "fonti": ["ANNCSU – Agenzia delle Entrate e ISTAT (CC BY 4.0)", "OpenStreetMap (ODbL)"],
}


INDICATORI = {
    "ente": {"nome": "Azienda USL della Romagna", "codice": "114", "regione_ce": "080", "fonte_residenti": "siseps",
             "anni": [{"anno": 2024, "provvisorio": False, "spesa": 39455736, "AZ9999": 2990000000, "BZ9999": 2966000000,
                       "leva": 75.78, "incidenza": 0.013196, "peso": 0.013303, "per_giornata": 45.8159,
                       "per_posto": 13438.6, "per_cittadino": 35.09},
                      {"anno": 2025, "provvisorio": True, "spesa": 40000000, "leva": None, "incidenza": None}]},
    "regione": {"nome": "Emilia-Romagna (consolidato)", "codice": "999", "regione_ce": "080",
                "anni": [{"anno": 2024, "provvisorio": False, "spesa": 187362060, "leva": 64.5, "incidenza": 0.0155}]},
    "italia": {"nome": "Italia", "codice": "IT", "anni": [{"anno": 2024, "provvisorio": False, "spesa": 2267333735,
                                                          "per_cittadino": 38.47}]},
    "componente": "totale energia e utenze", "anno_confronto": 2024,
    "fonti": [{"testo": "ISTAT – Popolazione residente (CC BY 4.0)", "url": "https://www.istat.it/note-legali/"},
              {"testo": "PUN GME", "url": "https://www.mercatoelettrico.org"}],
    "elaborazione": "Elaborazione di EuGEnio (Stato Critico) sui dati originali.",
}


def test_report_previsione():
    doc = RS.report_previsione(PREVISIONE)
    assert R.nome_file(doc, "xlsx") == "previsione_037006_breve.xlsx"
    t = leggi_pdf(R.genera(doc, "pdf"))
    for atteso in ("Bologna", "P50 (mediana)", "27/10/2026 00:00", "Weather data by Open-Meteo.com (CC BY 4.0)",
                   "TiRex-2", "ECMWF IFS", "4,2 %", "kWh"):
        assert atteso in t, atteso
    wb = libro_xlsx(R.genera(doc, "xlsx"))
    ws = wb["Previsione (kWh)"]
    assert ws["C2"].value == pytest.approx(9.4)                 # mediana numerica
    assert "Valore reale" in [c.value for c in ws[1]]
    assert "Bologna" in leggi_docx(R.genera(doc, "docx"))


def test_report_previsione_senza_richiesta():
    d = {k: v for k, v in PREVISIONE.items() if k not in ("richiesta", "verifica")}
    doc = RS.report_previsione(d, nome_comune=lambda i: "Bologna" if i == "037006" else None)
    assert doc.sottotitolo.startswith("Bologna")
    R.genera(doc, "docx")
    assert RS.report_previsione({}).titolo                      # risposta vuota: nessun errore


def test_report_legge10():
    doc = RS.report_legge10(LEGGE10)
    assert R.nome_file(doc, "pdf") == "clima_legge10_037006.pdf"
    t = leggi_pdf(R.genera(doc, "pdf"))
    for atteso in ("Bologna", "2.071", "-3,8 °C", "UNI 10349", "DPR 412/93", "Weather data by Open-Meteo.com",
                   "ISTAT (CC BY 4.0)", "Irradiazione"):
        assert atteso in t, atteso
    wb = libro_xlsx(R.genera(doc, "xlsx"))
    assert "Valori medi mensili" not in wb.sheetnames            # tabella senza titolo: nome generico
    assert any(ws["A2"].value == "gen" for ws in wb.worksheets)
    assert "non sostituiscono i valori normativi" in leggi_docx(R.genera(doc, "docx"))
    assert not any(x.titolo == "Confronto con il DPR 412/93" for x in doc.sezioni)   # senza dpr412 nel corpo


def test_report_legge10_dpr412():
    d = {**LEGGE10, "dpr412": {"gg": 2259, "zona": "E", "alt": 54, "nome": "", "nota": ""}}
    doc = RS.report_legge10(d)
    sez = next(x for x in doc.sezioni if x.titolo == "Confronto con il DPR 412/93")
    voci = {v.chiave: v.valore for b in sez.blocchi if hasattr(b, "voci") for v in b.voci}
    assert voci["Gradi giorno del DPR 412/93"] == 2259 and voci["Zona climatica del DPR 412/93"] == "E"
    assert voci["Differenza dei gradi giorno calcolati rispetto al DPR"] == pytest.approx(-8.3)   # 2071,4 vs 2259
    assert voci["La zona calcolata coincide con quella del DPR"] == "sì"
    t = leggi_pdf(R.genera(doc, "pdf"))
    for atteso in ("Confronto con il DPR 412/93", "2.259", "-8,3 %", "Normattiva", "valori di legge"):
        assert atteso in t, atteso
    assert "Confronto con il DPR" in leggi_docx(R.genera(doc, "docx"))
    R.genera(doc, "xlsx")
    # fusione: nessun valore unico, tabella dei comuni d'origine; zona diversa → "no"
    f = {**LEGGE10, "dpr412": {"gg": None, "zona": "D", "alt": None, "nota": "fusione di A, B",
                               "origini": [{"nome": "A", "gg": 2000, "zona": "D", "alt": 10},
                                           {"nome": "B", "gg": 2050, "zona": "D", "alt": 90}]}}
    sez = next(x for x in RS.report_legge10(f).sezioni if x.titolo == "Confronto con il DPR 412/93")
    tab = next(b for b in sez.blocchi if hasattr(b, "righe"))
    assert [r[1] for r in tab.righe] == [2000, 2050]
    voci = {v.chiave: v.valore for b in sez.blocchi if hasattr(b, "voci") for v in b.voci}
    assert voci["La zona calcolata coincide con quella del DPR"].startswith("no")
    assert "Differenza dei gradi giorno calcolati rispetto al DPR" not in voci
    leggi_pdf(R.genera(RS.report_legge10(f), "pdf"))
    # valori strani: nessun errore
    RS.report_legge10({**LEGGE10, "dpr412": {"gg": "x", "zona": "Q", "origini": "no"}})
    RS.report_legge10({**LEGGE10, "dpr412": [1, 2]})


def test_report_patrimonio_senza_dati_di_mercato():
    doc = RS.report_patrimonio(PATRIMONIO)
    assert R.nome_file(doc, "docx") == "patrimonio_geolocalizzato.docx"
    for formato, leggi in (("pdf", leggi_pdf), ("docx", leggi_docx)):
        t = leggi(R.genera(doc, formato))
        for atteso in ("Dati elaborati da Zornade", "https://zornade.com", "ANNCSU – Agenzia delle Entrate e ISTAT",
                       "ISTAT (CC BY 4.0)", "Pascoli", "Palestra", "da verificare", "1.235"):
            assert atteso in t, (formato, atteso)
        assert "prezzo_pun" not in t and "0,12" not in t
    wb = libro_xlsx(R.genera(doc, "xlsx"))
    ws = wb[wb.sheetnames[1]]
    intest = [c.value for c in ws[1]]
    assert "Latitudine" in intest and "prezzo_pun" not in intest
    assert ws.cell(row=2, column=intest.index("Latitudine") + 1).value == 44.4985


def test_report_patrimonio_colonne_sconosciute():
    doc = RS.report_patrimonio({"righe": [{"codice_x": "A", "valore_y": 3, "psv_medio": 40}]})
    t = leggi_docx(R.genera(doc, "docx"))
    assert "codice_x" in t and "psv_medio" not in t


# ---------------------------------------------------------------------------------------------- rotte


@pytest.fixture(scope="module")
def client():
    """app.py senza TiRex-2 né fase 3: i moduli meteo_storico sono segnaposto vuoti."""
    from fastapi.testclient import TestClient
    finti = {"meteo_storico": {"VARIABILI": []}, "meteo_storico.aggrega": {"giornaliero": None},
             "meteo_storico.correzione": {"applica": None}, "meteo_storico.limiti": {"Limitatore": object},
             "meteo_storico.scarica": {"FONTI": {}, "client": None, "richiesta": None}}
    salvati = {k: sys.modules.get(k) for k in finti}
    for nome, attr in finti.items():
        if nome not in sys.modules:
            m = types.ModuleType(nome)
            m.__dict__.update(attr)
            sys.modules[nome] = m
    from servizio import app as A
    A._storia.clear()
    yield TestClient(A.app), A
    for k, v in salvati.items():
        if v is None:
            sys.modules.pop(k, None)


def test_report_indicatori():
    doc = RS.report_indicatori(INDICATORI)
    assert R.nome_file(doc, "pdf") == "indicatori_energia_080_114.pdf"
    t = leggi_pdf(R.genera(doc, "pdf"))
    for atteso in ("Azienda USL della Romagna", "39.455.736", "75,8", "1,32", "45,82", "Confronto nel 2024",
                   "2025 (provvisorio)", "OpenBDAP", "CC BY 4.0", "ISTAT"):
        assert atteso in t, atteso
    assert "PUN" not in t                                        # nessun dato di mercato, nemmeno tra le fonti
    assert libro_xlsx(R.genera(doc, "xlsx")).worksheets
    assert "Avvertenze" in leggi_docx(R.genera(doc, "docx"))
    assert RS.report_indicatori({}).titolo                       # corpo vuoto: nessun errore


def test_rotta_generica(client):
    c, _ = client
    assert R.formati_attivi() == ("pdf", "xlsx", "ods")             # dal 28/10/2026: solo PDF e dati grezzi
    for formato in R.formati_attivi():
        tipo = R.FORMATI[formato]
        r = c.post(f"/api/report?formato={formato}", json=ESEMPIO)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"] == tipo
        assert r.headers["content-disposition"] == f'attachment; filename="relazione_di_prova_forli_20261027.{formato}"'
    assert "Forlì" in leggi_pdf(c.post("/api/report?formato=pdf", json=ESEMPIO).content)


def test_rotte_servizi(client):
    c, _ = client
    casi = {"previsione": (PREVISIONE, "previsione_037006_breve"), "legge10": (LEGGE10, "clima_legge10_037006"),
            "patrimonio": (PATRIMONIO, "patrimonio_geolocalizzato"),
            "indicatori": (INDICATORI, "indicatori_energia_080_114")}
    for tipo, (corpo, nome) in casi.items():
        for formato in R.formati_attivi():
            r = c.post(f"/api/report/{tipo}?formato={formato}", json=corpo)
            assert r.status_code == 200, (tipo, formato, r.text)
            assert f'filename="{nome}.{formato}"' in r.headers["content-disposition"]


def test_rotte_errori(client, monkeypatch):
    c, A = client
    r = c.post("/api/report?formato=rtf", json=ESEMPIO)
    assert r.status_code == 400 and "Formato non supportato" in r.json()["detail"]
    r = c.post("/api/report/bolletta?formato=pdf", json={})
    assert r.status_code == 400 and "Tipo di report sconosciuto" in r.json()["detail"]
    r = c.post("/api/report?formato=pdf", json={"sezioni": []})
    assert r.status_code == 422 and "titolo: campo obbligatorio" in r.json()["detail"]
    r = c.post("/api/report/legge10?formato=pdf", content=b"non json", headers={"content-type": "application/json"})
    assert r.status_code == 400
    r = c.post("/api/report/legge10?formato=pdf", json=[1, 2])
    assert r.status_code == 422
    monkeypatch.setattr(A.C, "REPORT_MAX_MB", 0.001)
    r = c.post("/api/report?formato=pdf", json=ESEMPIO)
    assert r.status_code == 413 and "troppo grande" in r.json()["detail"]
    monkeypatch.setattr(A.C, "REPORT_MAX_MB", 5)
    monkeypatch.setitem(A.LIMITI, "report", (1, 3600))
    A._storia.clear()
    assert c.post("/api/report?formato=xlsx", json=ESEMPIO).status_code == 200
    r = c.post("/api/report?formato=xlsx", json=ESEMPIO)
    assert r.status_code == 429 and "Troppe richieste" in r.json()["detail"]
    A._storia.clear()


def test_word_e_odt_spenti(client, monkeypatch):
    """Word e ODT sono spenti: 410 con un messaggio che indica PDF e dati grezzi; si riaccendono con REPORT_FORMATI."""
    c, A = client
    A._storia.clear()
    for formato in ("docx", "odt", "DOCX"):
        for url in (f"/api/report?formato={formato}", f"/api/report/legge10?formato={formato}"):
            r = c.post(url, json=ESEMPIO)
            assert r.status_code == 410, (url, r.text)
            det = r.json()["detail"]
            assert "non è più disponibile" in det and "solo in PDF" in det and "Excel o ODS" in det
    assert c.get("/api/report/schema").status_code == 200
    monkeypatch.setattr(A.C, "REPORT_FORMATI", ("pdf", "xlsx", "ods", "docx", "odt"))
    for formato in ("docx", "odt"):
        r = c.post(f"/api/report?formato={formato}", json=ESEMPIO)
        assert r.status_code == 200 and r.headers["content-type"] == R.FORMATI[formato]
    monkeypatch.setattr(A.C, "REPORT_FORMATI", ("pdf", "rtf"))      # valori sconosciuti nel .env: ignorati
    assert R.formati_attivi() == ("pdf",)
    A._storia.clear()


def test_schema(client):
    c, _ = client
    s = c.get("/api/report/schema").json()
    assert "sezioni" in s["properties"]


def test_legge10_senza_archivio_meteo(client, monkeypatch):
    """Senza l'archivio della fase 3 sul server: 503 con un messaggio chiaro, non "Errore interno"."""
    c, A = client
    def manca(istat):
        raise FileNotFoundError("giornaliero.parquet")
    monkeypatch.setattr(A.meteo, "storico_giornaliero", manca)
    monkeypatch.setattr(A, "_istat_valido", lambda istat: None)
    finto = types.ModuleType("meteo_storico.legge10")
    finto.calcola = lambda d, anni: {}
    monkeypatch.setitem(sys.modules, "meteo_storico.legge10", finto)
    monkeypatch.setattr(sys.modules["meteo_storico"], "legge10", finto, raising=False)
    r = c.get("/api/clima-legge10/037006")
    assert r.status_code == 503 and "non sono ancora sul server" in r.json()["detail"]


# ---------------------------------------------------------------------------------------------- valori per sede
def _svg_mappa(n_comuni=40, testo="Consumi per sede – 1 mattoncino = 2 GWh · 250.000 €") -> str:
    """SVG piatto come quello che manda il sito: percorsi, poligoni delle barre, tratteggi in px, testo con € e –."""
    parti = ['<?xml version="1.0" encoding="UTF-8"?>\n<svg xmlns="http://www.w3.org/2000/svg" width="600" height="400" '
             'viewBox="0 0 600 400"><title>Mappa</title><rect x="0" y="0" width="600" height="400" fill="rgb(247,246,248)"/>']
    for i in range(n_comuni):
        x, y = 20 + (i % 10) * 50, 40 + (i // 10) * 60
        parti.append(f'<path d="M{x},{y}L{x + 45},{y + 5}L{x + 40},{y + 50}L{x + 3},{y + 45}Z" fill="rgb(225,221,230)" '
                     f'stroke="rgb(247,246,248)" stroke-width="0.5px"/>')
    parti.append('<polygon points="300,300 320,290 320,150 300,160" fill="rgb(196,60,18)" stroke="rgb(90,20,5)" stroke-width="0.5px"/>')
    parti.append('<line x1="10" y1="20" x2="590" y2="20" stroke="rgb(63,75,208)" stroke-dasharray="7px, 4px" stroke-width="1.5px"/>')
    parti.append(f'<text x="12" y="390" font-family="Hanken Grotesk, Helvetica, Arial, sans-serif" font-size="12px" '
                 f'fill="rgb(28,24,32)">{testo}</text></svg>')
    return base64.b64encode("".join(parti).encode()).decode()


def corpo_valori(n_mappe=2, n_pareto=6, n_comuni=40):
    sedi = [f"H{i:02d} · Ospedale {i}" for i in range(1, 14)]
    return {
        "titolo": "Valori per sede", "sottotitolo": "13 sedi – Emilia-Romagna",
        "totali": [{"colonna": "Energia elettrica 2025", "tipo": "consumo", "totale": 85637.8, "unita": "MWh", "sedi": 13},
                   {"colonna": "Gas 2025", "tipo": "consumo", "totale": 20984431, "unita": "Smc", "sedi": 13},
                   {"colonna": "Costo energia elettrica 2025", "tipo": "costo", "totale": 25895164.8, "unita": "€", "sedi": 13},
                   {"colonna": "PUN medio", "tipo": "costo", "totale": 120.5, "unita": "€/MWh", "sedi": 13}],
        "prezzi": [{"consumo": "Energia elettrica 2025", "costo": "Costo energia elettrica 2025", "valore": 302.38,
                    "unita": "€/MWh", "sedi": 13}],
        "mappe": [{"titolo": f"Mappa {i}", "svg": _svg_mappa(n_comuni)} for i in range(n_mappe)],
        "pareto": [{"titolo": f"Pareto: colonna {i}", "tipo": "consumo" if i % 2 == 0 else "costo", "svg": _svg_mappa(5, "80 %")}
                   for i in range(n_pareto)],
        "tabelle": [{"titolo": "Consumi per sede", "tipo": "consumo", "intestazioni": ["Sede", "Comune", "Energia (MWh)", "PSV"],
                     "righe": [[s, "Ravenna", 1000.5 + i, 7] for i, s in enumerate(sedi)], "decimali": [None, None, 0, 2]},
                    {"titolo": "Costi per sede", "tipo": "costo", "intestazioni": ["Sede", "Comune", "Costo (€)"],
                     "righe": [[s, "Ravenna", 300000 + i] for i, s in enumerate(sedi)], "decimali": [None, None, 0]}],
        "note": ["Nota di prova."], "fonti": ["Numeri civici: ANNCSU – Agenzia delle Entrate e ISTAT (CC BY 4.0)"],
        "avviso": {"titolo": "Dati di esempio non reali", "testo": "Consumi e costi sono numeri inventati."},
    }


def test_report_valori_sedi_pdf_ordine_e_immagini():
    doc = RS.report_valori_sedi(corpo_valori(), formato="pdf")
    titoli = [s.titolo for s in doc.sezioni]
    attesi = ["Totali", "Costi unitari medi", "Mappe tridimensionali", "Diagrammi di Pareto dei consumi",
              "Diagrammi di Pareto dei costi", "Valori per sede: consumi", "Valori per sede: costi"]
    assert [t for t in titoli if t in attesi] == attesi          # l'ordine chiesto da Paolo
    immagini = [b for s in doc.sezioni for b in s.blocchi if isinstance(b, R.Immagine)]
    assert len(immagini) == 8 and all(i.formato == "svg" for i in immagini)
    pdf = R.genera(doc, "pdf")
    t = leggi_pdf(pdf)
    for atteso in ("Dati di esempio non reali", "85.638", "20.984.431", "302,38", "€/MWh", "Costi unitari medi", "Abbinamento",
                   "Dati elaborati da Zornade", "Ospedale 13"):
        assert atteso in t, atteso
    assert "[Immagine non riproducibile" not in t                # fpdf2 legge i nostri SVG (px, €, tratteggi)
    assert "PUN" not in t and "PSV" not in t                      # nessun dato di mercato, nemmeno per nome di colonna
    assert t.count("ANNCSU") == 1                                 # la fonte ANNCSU arrivata dal sito non si ripete


def test_report_valori_sedi_fogli_senza_immagini():
    doc = RS.report_valori_sedi(corpo_valori(), formato="xlsx")
    assert not any(isinstance(b, R.Immagine) for s in doc.sezioni for b in s.blocchi)
    wb = libro_xlsx(R.genera(doc, "xlsx"))
    nomi = [ws.title for ws in wb.worksheets]
    assert "Consumi per sede" in nomi and "Costi per sede" in nomi and "Costi unitari medi" in nomi
    testo = "\n".join(str(c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r if c.value is not None)
    assert "Immagine non inclusa" not in testo and "PUN" not in testo and "PSV" not in testo
    assert wb["Consumi per sede"]["C2"].value == 1000.5           # i numeri restano numeri
    assert R.genera(RS.report_valori_sedi(corpo_valori(), formato="ods"), "ods")[:2] == b"PK"


def test_report_valori_sedi_difensivo():
    assert RS.report_valori_sedi({}).titolo == "Valori per sede"            # corpo vuoto: nessun errore
    c = corpo_valori(n_mappe=1, n_pareto=0)
    c["mappe"].append({"titolo": "rotta", "svg": "non base64 !"})
    c["mappe"].append({"titolo": "esterna", "svg": base64.b64encode(b'<svg xmlns="http://www.w3.org/2000/svg"><image href="http://x/y.png"/></svg>').decode()})
    doc = RS.report_valori_sedi(c)
    assert sum(isinstance(b, R.Immagine) for s in doc.sezioni for b in s.blocchi) == 1   # le immagini non valide si saltano
    molte = RS.report_valori_sedi(corpo_valori(n_mappe=2, n_pareto=12))
    assert sum(isinstance(b, R.Immagine) for s in molte.sezioni for b in s.blocchi) == R.C.REPORT_MAX_IMMAGINI
    assert "non sono entrati" in leggi_pdf(R.genera(molte, "pdf"))


def test_svg_per_pdf_senza_dejavu():
    s = R._svg_per_pdf(base64.b64decode(_svg_mappa(1)), "helvetica", False).decode()
    assert 'font-family="helvetica"' in s and "<title>" not in s and "EUR" in s and "px" not in s.split("<text")[0].split("stroke-dasharray")[1][:12]


def test_rotta_valori_sedi_dimensioni(client):
    """2 mappe grandi (come una regione intera, ~300 comuni) + 6 Pareto stanno nei limiti attuali (REPORT_MAX_MB)."""
    c, _ = client
    import json
    corpo = json.dumps(corpo_valori(n_mappe=2, n_pareto=6, n_comuni=1500)).encode()
    assert len(corpo) < R.C.REPORT_MAX_MB * 1024 * 1024
    r = c.post("/api/report/valori_sedi?formato=pdf", content=corpo, headers={"Content-Type": "application/json"})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert "valori_per_sede" in r.headers["content-disposition"]
    r = c.post("/api/report/valori_sedi?formato=xlsx", content=corpo, headers={"Content-Type": "application/json"})
    assert r.status_code == 200 and r.content[:2] == b"PK"


def test_report_valori_sedi_decimali_dei_prezzi():
    """Elettricità in €/MWh con 2 decimali, gas e acqua con i decimali indicati dalla pagina (o 3)."""
    c = corpo_valori()
    c["prezzi"] = [{"consumo": "Energia elettrica 2025", "costo": "Costo energia elettrica 2025", "valore": 302.384,
                    "unita": "€/MWh", "sedi": 13, "decimali": 2},
                   {"consumo": "Gas 2025", "costo": "Costo gas 2025", "valore": 1.03149, "unita": "€/Smc", "sedi": 13}]
    t = leggi_pdf(R.genera(RS.report_valori_sedi(c, formato="pdf"), "pdf"))
    assert "302,38" in t and "302,384" not in t and "1,031" in t
