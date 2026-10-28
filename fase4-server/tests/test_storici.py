"""Prove della pagina "Dati storici" (servizio/storici.py) su una cartella dati sintetica generata qui.

  cd fase4-server && python -m pytest tests/test_storici.py -q
"""
from __future__ import annotations

import datetime as dt
import io
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

RADICE = Path(__file__).resolve().parents[2]
for p in (RADICE / "fase4-server", RADICE / "fase3-storico"):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from servizio import config as C, storici  # noqa: E402

COMUNI = [  # istat, nome, belfiore, provincia, regione_cod, regione, capoluogo
    ("033001", "Agazzano", "A067", "PC", "08", "Emilia-Romagna", False),
    ("033032", "Piacenza", "G535", "PC", "08", "Emilia-Romagna", True),
    ("037006", "Bologna", "A944", "BO", "08", "Emilia-Romagna", True),
    ("048017", "Firenze", "D612", "FI", "09", "Toscana", True),
    ("027042", "Venezia", "L736", "VE", "05", "Veneto", True),     # fuori da REGIONI: il server non lo serve
]
ORARIO_DAL, ORARIO_GIORNI = pd.Timestamp("2023-01-01", tz="UTC"), 21


def _giornaliero() -> pd.DataFrame:
    date = pd.date_range("2019-07-01", "2024-06-30", freq="D")
    parti = []
    for k, (istat, *_ ) in enumerate(COMUNI):
        doy = date.dayofyear.to_numpy()
        t = 13 + 11 * np.sin(2 * np.pi * (doy - 110) / 365.25) + k
        d = pd.DataFrame({"istat": istat, "data": date.date, "t_media": t, "t_min": t - 5, "t_max": t + 6,
                          "ur_media": 70.0, "rad_globale_kwh_m2": 3.0, "rad_diffusa_kwh_m2": 1.2,
                          "vento_medio_kmh": 8.0, "nuvolosita_media": 50.0, "pioggia_mm": 1.0, "neve_cm": 0.0})
        if istat == "033001":   # gennaio 2021 con 10 giorni mancanti
            d = d[~((date >= "2021-01-05") & (date <= "2021-01-14"))]
        d["gg20"] = (20 - d.t_media).clip(lower=0)
        d["gc24"] = (d.t_media - 24).clip(lower=0)
        parti.append(d)
    g = pd.concat(parti, ignore_index=True)
    for c in g.columns[2:]:
        g[c] = g[c].astype("float32")
    return g


def _orario(k: int) -> pd.DataFrame:
    t = pd.date_range(ORARIO_DAL, periods=24 * ORARIO_GIORNI, freq="h")
    ora = t.hour.to_numpy()
    return pd.DataFrame({
        "time": t, "temperature_2m": (ora + k).astype("float32"), "relative_humidity_2m": np.float32(80),
        "shortwave_radiation": np.where(ora == 11, 1000.0, 0.0).astype("float32"),
        "diffuse_radiation": np.where(ora == 11, 400.0, 0.0).astype("float32"),
        "wind_speed_10m": np.float32(10), "cloud_cover": np.float32(25),
        "precipitation": np.float32(0.5), "snowfall": np.float32(0)})


@pytest.fixture(scope="session")
def cartella(tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("dati")
    cm = pd.DataFrame(COMUNI, columns=["istat", "nome", "belfiore", "provincia", "regione_cod", "regione", "capoluogo"])
    cm["lat"], cm["lon"], cm["metodo_punto"] = 44.5, 11.3, "edifici_2021"
    cm.to_parquet(d / "comuni.parquet", index=False)
    _giornaliero().to_parquet(d / "giornaliero.parquet", index=False, compression="zstd")
    (d / "orario").mkdir()
    for k, (istat, *_ ) in enumerate(COMUNI[:4]):
        if istat != "037006":      # Bologna senza file orario
            _orario(k).to_parquet(d / "orario" / f"{istat}.parquet", index=False)
    e = d / "mercati" / "eurostat"
    e.mkdir(parents=True)
    righe = []
    for sem, v in (("2023-S1", 0.25), ("2023-S2", 0.20), ("2024-S1", 0.18)):
        for tax in ("X_TAX", "X_VAT", "I_TAX"):
            for fascia, nome in (("MWH20-499", "20 MWh - 499 MWh"), ("MWH500-1999", "500 MWh - 1 999 MWh")):
                righe.append({"serie": "elettricita_pmi", "freq": "S", "product": "6000", "nrg_cons": fascia,
                              "unit": "KWH", "tax": tax, "currency": "EUR", "geo": "IT", "time": sem,
                              "valore": v, "fascia_consumo": nome})
    pd.DataFrame(righe).to_parquet(e / "elettricita_pmi.parquet", index=False)
    # dati interni che non devono mai uscire
    (d / "mercati" / "mgp_prezzi").mkdir()
    pd.DataFrame({"data": [dt.date(2024, 1, 1)], "pun": [100.0]}).to_parquet(d / "mercati" / "mgp_prezzi" / "2024.parquet")
    (d / "anncsu").mkdir()
    pd.DataFrame({"via": ["x"]}).to_parquet(d / "anncsu" / "anncsu_civici.parquet")
    return d


@pytest.fixture()
def cliente(cartella, monkeypatch) -> TestClient:
    monkeypatch.setattr(C, "DATI", cartella)
    monkeypatch.setattr(C, "REGIONI", ["08", "09"])
    monkeypatch.setattr(C, "SOLO_COMUNI", [])
    storici._comuni_cache.cache_clear()
    storici._disponibilita.cache_clear()
    storici._storia.clear()
    app = FastAPI()
    app.include_router(storici.router)
    return TestClient(app)


def _leggi_csv(testo: bytes) -> pd.DataFrame:
    assert testo.startswith("﻿".encode("utf-8"))
    return pd.read_csv(io.BytesIO(testo), sep=";", decimal=",", comment="#", encoding="utf-8-sig",
                       dtype={"istat": str})


def _estrai(cl, corpo, formato="csv"):
    r = cl.post(f"/api/storici/estrai?formato={formato}", json=corpo)
    assert r.status_code == 200, r.text
    return r


# ---------------------------------------------------------------------------------------------- catalogo
def test_catalogo(cliente):
    d = cliente.get("/api/storici/catalogo").json()
    ids = [x["id"] for x in d["dataset"]]
    assert ids == ["meteo_orario", "meteo_giornaliero", "gradi_giorno", "legge10", "comuni", "eurostat_pmi"]
    assert d["regioni"] == ["08", "09"] and d["comuni"] == 4
    g = {x["id"]: x for x in d["dataset"]}
    assert g["meteo_giornaliero"]["inizio"] == "2019-07-01" and g["meteo_giornaliero"]["fine"] == "2024-06-30"
    assert g["meteo_orario"]["inizio"] == "2023-01-01"
    assert g["eurostat_pmi"]["inizio"] == "2023-01-01" and g["eurostat_pmi"]["fine"] == "2024-06-30"
    assert [v["id"] for v in g["eurostat_pmi"]["variabili"]] == ["elettricita_pmi"]
    assert "Open-Meteo" in g["meteo_orario"]["fonte"] and g["meteo_orario"]["licenza"] == "CC BY 4.0"
    testo = str(d).lower()
    assert "mgp_prezzi" not in testo and "anncsu_civici" not in testo
    assert any("GME" in x["nome"] for x in d["non_scaricabili"])


def test_catalogo_solo_file_presenti(tmp_path, cartella, monkeypatch):
    import shutil
    shutil.copy(cartella / "comuni.parquet", tmp_path / "comuni.parquet")
    shutil.copy(cartella / "giornaliero.parquet", tmp_path / "giornaliero.parquet")
    monkeypatch.setattr(C, "DATI", tmp_path)
    monkeypatch.setattr(C, "REGIONI", ["08"])
    storici._comuni_cache.cache_clear()
    storici._disponibilita.cache_clear()
    app = FastAPI()
    app.include_router(storici.router)
    ids = [x["id"] for x in TestClient(app).get("/api/storici/catalogo").json()["dataset"]]
    assert ids == ["meteo_giornaliero", "gradi_giorno", "legge10", "comuni"]
    r = TestClient(app).post("/api/storici/stima", json={"dataset": "meteo_orario", "area": {"tipo": "tutti"}})
    assert r.status_code == 404


# ---------------------------------------------------------------------------------------------- aggregazioni
def test_mensile_giornaliero(cliente):
    corpo = {"dataset": "meteo_giornaliero", "frequenza": "mensile", "dal": "2020-01-10", "al": "2021-02-03",
             "variabili": ["t_media", "t_min", "t_max", "pioggia_mm", "rad_globale_kwh_m2"],
             "area": {"tipo": "comuni", "comuni": ["033001", "037006"]}}
    s = cliente.post("/api/storici/stima", json=corpo).json()
    assert s["consentito"] and s["righe"] == 2 * 14          # da gennaio 2020 a febbraio 2021, mesi interi
    assert s["dal"] == "2020-01-01" and s["al"] == "2021-02-28"
    df = _leggi_csv(_estrai(cliente, corpo).content)
    assert list(df.columns[:6]) == ["istat", "comune", "provincia", "periodo", "dati_validi (n)", "dati_attesi (n)"]
    assert "t_media (°C)" in df.columns and "pioggia (mm)" not in df.columns and "precipitazione (mm)" in df.columns
    g = _giornaliero()
    g["data"] = pd.to_datetime(g.data)
    rif = g[(g.istat == "037006") & (g.data.dt.strftime("%Y-%m") == "2020-03")]
    r = df[(df.istat == "037006") & (df.periodo == "2020-03")].iloc[0]
    assert r["dati_attesi (n)"] == 31 and r["dati_validi (n)"] == 31
    assert r["t_media (°C)"] == pytest.approx(rif.t_media.astype(float).mean(), abs=0.01)
    assert r["t_min (°C)"] == pytest.approx(rif.t_min.astype(float).min(), abs=0.01)
    assert r["t_max (°C)"] == pytest.approx(rif.t_max.astype(float).max(), abs=0.01)
    assert r["precipitazione (mm)"] == pytest.approx(31.0)
    assert r["rad_globale (kWh/m²)"] == pytest.approx(93.0)
    # gennaio 2021 di Agazzano: 21 giorni su 31, sotto il 90% → vuoto; con copertura 50% c'è
    gen = df[(df.istat == "033001") & (df.periodo == "2021-01")].iloc[0]
    assert gen["dati_validi (n)"] == 21 and pd.isna(gen["t_media (°C)"])
    df50 = _leggi_csv(_estrai(cliente, {**corpo, "copertura": 50}).content)
    assert not pd.isna(df50[(df50.istat == "033001") & (df50.periodo == "2021-01")]["t_media (°C)"].iloc[0])


def test_annuale_e_stagione_gradi_giorno(cliente):
    g = _giornaliero()
    g["data"] = pd.to_datetime(g.data)
    x = g[g.istat == "033032"]
    base = {"dataset": "gradi_giorno", "variabili": ["gg20", "gc24"], "area": {"tipo": "comuni", "comuni": ["033032"]}}
    an = _leggi_csv(_estrai(cliente, {**base, "frequenza": "annuale"}).content)
    assert list(an.periodo.astype(str)) == ["2019", "2020", "2021", "2022", "2023", "2024"]
    r2020 = an[an.periodo.astype(str) == "2020"].iloc[0]
    assert r2020["dati_attesi (n)"] == 366
    assert r2020["gg20 (°C·giorno)"] == pytest.approx(x[x.data.dt.year == 2020].gg20.astype(float).sum(), abs=0.1)
    assert pd.isna(an[an.periodo.astype(str) == "2019"]["gg20 (°C·giorno)"].iloc[0])   # solo luglio–dicembre
    st = _leggi_csv(_estrai(cliente, {**base, "frequenza": "stagione"}).content)
    assert list(st.periodo) == ["2019/20", "2020/21", "2021/22", "2022/23", "2023/24"]
    sel = x[(x.data >= "2021-07-01") & (x.data <= "2022-06-30")]
    r = st[st.periodo == "2021/22"].iloc[0]
    assert r["dati_attesi (n)"] == 365 and r["gg20 (°C·giorno)"] == pytest.approx(sel.gg20.astype(float).sum(), abs=0.1)
    gio = _leggi_csv(_estrai(cliente, {**base, "frequenza": "giornaliera", "dal": "2022-01-01", "al": "2022-01-31"}).content)
    assert len(gio) == 31 and "dati_validi (n)" not in gio.columns


def test_orario(cliente):
    corpo = {"dataset": "meteo_orario", "frequenza": "oraria", "dal": "2023-01-02", "al": "2023-01-04",
             "variabili": ["temperature_2m", "shortwave_radiation"], "area": {"tipo": "provincia", "provincia": "PC"}}
    s = cliente.post("/api/storici/stima", json=corpo).json()
    assert s["righe"] == 2 * 3 * 24
    df = _leggi_csv(_estrai(cliente, corpo).content)
    assert len(df) == 2 * 3 * 24
    assert df.columns.tolist() == ["istat", "comune", "provincia", "ora_utc (UTC)", "ora_locale (ora italiana)",
                                   "temperatura (°C)", "radiazione_globale (W/m²)"]
    prima = df[df.istat == "033001"].iloc[0]
    assert prima["ora_locale (ora italiana)"] == "2023-01-02 00:00" and prima["ora_utc (UTC)"] == "2023-01-01T23:00Z"
    # giornaliero dall'orario: radiazione in kWh/m², minimi e massimi orari
    gio = _leggi_csv(_estrai(cliente, {**corpo, "frequenza": "giornaliera"}).content)
    r = gio[(gio.istat == "033032") & (gio.data == "2023-01-03")].iloc[0]
    assert r["radiazione_globale (kWh/m²)"] == pytest.approx(1.0)
    assert r["t_min (°C)"] == pytest.approx(1.0) and r["t_max (°C)"] == pytest.approx(24.0)  # ora locale = UTC+1
    assert r["dati_attesi (n)"] == 24
    # il mese di gennaio 2023 ha solo 21 giorni di dati orari: vuoto al 90%, presente al 60%
    men = _leggi_csv(_estrai(cliente, {**corpo, "frequenza": "mensile"}).content)
    assert pd.isna(men["t_media (°C)"]).all() and (men["dati_attesi (n)"] == 744).all()
    men = _leggi_csv(_estrai(cliente, {**corpo, "frequenza": "mensile", "copertura": 60}).content)
    assert men["radiazione_globale (kWh/m²)"].tolist() == pytest.approx([21.0, 21.0], abs=0.05)


def test_frequenza_non_valida(cliente):
    r = cliente.post("/api/storici/stima", json={"dataset": "meteo_giornaliero", "frequenza": "oraria",
                                                 "area": {"tipo": "regione", "regione": "08"}})
    assert r.status_code == 422


# ---------------------------------------------------------------------------------------------- area
@pytest.mark.parametrize("area,attesi", [
    ({"tipo": "regione", "regione": "08"}, {"033001", "033032", "037006"}),
    ({"tipo": "regione", "regione": "09"}, {"048017"}),
    ({"tipo": "provincia", "provincia": "pc"}, {"033001", "033032"}),
    ({"tipo": "regione", "regione": "08", "solo_capoluoghi": True}, {"033032", "037006"}),
    ({"tipo": "tutti"}, {"033001", "033032", "037006", "048017"}),
    ({"tipo": "comuni", "comuni": ["048017", "033001"]}, {"048017", "033001"}),
])
def test_filtri_geografici(cliente, area, attesi):
    df = _leggi_csv(_estrai(cliente, {"dataset": "comuni", "area": area}).content)
    assert set(df.istat) == attesi


def test_aree_non_valide(cliente):
    stima = lambda a: cliente.post("/api/storici/stima", json={"dataset": "comuni", "area": a})  # noqa: E731
    assert stima({"tipo": "regione", "regione": "05"}).status_code == 422          # regione non servita
    assert stima({"tipo": "comuni", "comuni": []}).status_code == 422
    assert stima({"tipo": "comuni", "comuni": ["12ab56"]}).status_code == 422
    r = stima({"tipo": "comuni", "comuni": ["033001", "027042"]}).json()        # Venezia esclusa con avviso
    assert r["comuni"] == 1 and any("027042" in a for a in r["avvisi"])
    assert stima({"tipo": "provincia", "provincia": "FI", "solo_capoluoghi": True}).json()["comuni"] == 1


def test_elenco_comuni(cliente):
    df = _leggi_csv(_estrai(cliente, {"dataset": "comuni", "area": {"tipo": "tutti"}}).content)
    assert {"istat", "comune", "belfiore", "provincia", "capoluogo", "lat (°)", "lon (°)"} <= set(df.columns)
    assert df.set_index("istat").capoluogo.to_dict()["033001"] == "no"
    assert "027042" not in set(df.istat)


# ---------------------------------------------------------------------------------------------- Legge 10
def test_legge10(cliente):
    base = {"dataset": "legge10", "stagioni": 3, "area": {"tipo": "provincia", "provincia": "PC"}}
    men = _leggi_csv(_estrai(cliente, {**base, "tipologia": "mensili"}).content)
    assert len(men) == 24 and sorted(men.mese.unique()) == list(range(1, 13))
    assert men.stagioni.iloc[0] == "2021/22 – 2023/24"
    rie = _leggi_csv(_estrai(cliente, {**base, "tipologia": "riepilogo"}).content)
    assert len(rie) == 2 and set(rie.zona_risultante) <= set("ABCDEF")
    assert (rie["n_stagioni (n)"] == 3).all()
    st = _leggi_csv(_estrai(cliente, {**base, "tipologia": "stagioni"}).content)
    assert len(st) == 6 and set(st.stagione) == {"2021/22", "2022/23", "2023/24"}
    testo = _estrai(cliente, {**base, "tipologia": "riepilogo"}).content.decode("utf-8-sig")
    assert "UNI 10349" in testo and "Open-Meteo" in testo


def test_legge10_troppi_comuni(cliente, monkeypatch):
    monkeypatch.setattr(C, "STORICI_MAX_COMUNI_LEGGE10", 2)
    r = cliente.post("/api/storici/stima", json={"dataset": "legge10", "area": {"tipo": "tutti"}})
    assert r.status_code == 413


# ---------------------------------------------------------------------------------------------- Eurostat
def test_eurostat(cliente):
    corpo = {"dataset": "eurostat_pmi", "variabili": ["elettricita_pmi"], "tassazione": ["X_VAT"],
             "dal": "2023-07-01", "al": "2024-12-31"}
    df = _leggi_csv(_estrai(cliente, corpo).content)
    assert len(df) == 4 and set(df.tassazione) == {"X_VAT"} and set(df.semestre) == {"2023-S2", "2024-S1"}
    assert df["valore_eur_mwh (€/MWh)"].max() == pytest.approx(200.0)
    assert "Eurostat" in _estrai(cliente, corpo).content.decode("utf-8-sig")


# ---------------------------------------------------------------------------------------------- limiti
def test_limiti_righe(cliente, monkeypatch):
    monkeypatch.setattr(C, "STORICI_MAX_RIGHE_CSV", 1000)
    monkeypatch.setattr(C, "STORICI_MAX_RIGHE_XLSX", 500)
    corpo = {"dataset": "meteo_giornaliero", "frequenza": "giornaliera", "area": {"tipo": "tutti"}}
    s = cliente.post("/api/storici/stima", json=corpo).json()
    assert not s["consentito"] and s["righe"] > 1000 and s["suggerimenti"]
    assert any("mensile" in x for x in s["suggerimenti"])
    r = cliente.post("/api/storici/estrai?formato=csv", json=corpo)
    assert r.status_code == 413 and "Per ridurla" in r.json()["detail"]
    ok = {**corpo, "frequenza": "mensile"}                     # 4 comuni × 60 mesi = 240 righe
    assert cliente.post("/api/storici/stima?formato=xlsx", json=ok).json()["consentito"]
    s = cliente.post("/api/storici/stima?formato=xlsx", json={**corpo, "dal": "2023-01-01", "al": "2023-06-30"}).json()
    assert not s["consentito"] and any("CSV" in x for x in s["suggerimenti"])
    monkeypatch.setattr(C, "STORICI_MAX_RIGHE_ODS", 500)
    assert cliente.post("/api/storici/stima?formato=ods", json=ok).json()["consentito"]
    s = cliente.post("/api/storici/stima?formato=ods", json={**corpo, "dal": "2023-01-01", "al": "2023-06-30"}).json()
    assert not s["consentito"] and "OpenDocument" in s["motivo"] and any("CSV" in x for x in s["suggerimenti"])
    monkeypatch.setattr(C, "STORICI_MAX_CELLE_ODS", 100)
    s = cliente.post("/api/storici/stima?formato=ods", json=ok).json()
    assert not s["consentito"] and "celle" in s["motivo"]


def test_limite_per_ip(cliente, monkeypatch):
    monkeypatch.setattr(C, "STORICI_ESTRAZIONI_ORA", 2)
    corpo = {"dataset": "comuni", "area": {"tipo": "tutti"}}
    assert cliente.post("/api/storici/estrai", json=corpo).status_code == 200
    assert cliente.post("/api/storici/estrai", json=corpo).status_code == 200
    r = cliente.post("/api/storici/estrai", json=corpo)
    assert r.status_code == 429 and "all'ora" in r.json()["detail"]


@pytest.mark.parametrize("corpo", [
    {"dataset": "pun"}, {"dataset": "mgp_prezzi"}, {"dataset": "prezzi_zonali"}, {"dataset": "psv"},
    {"dataset": "igi"}, {"dataset": "ttf_forward"}, {"dataset": "eex_future"}, {"dataset": "pfor"},
    {"dataset": "psv_arera"}, {"dataset": "zornade"}, {"dataset": "anncsu"},
    {"dataset": "eurostat_pmi", "variabili": ["../mgp_prezzi/2024"]},
    {"dataset": "meteo_giornaliero", "variabili": ["t_media", "PUN"]},
])
def test_dataset_vietati(cliente, corpo):
    for rotta in ("/api/storici/stima", "/api/storici/estrai?formato=csv", "/api/storici/estrai?formato=xlsx",
                  "/api/storici/stima?formato=ods", "/api/storici/estrai?formato=ods"):
        r = cliente.post(rotta, json={"area": {"tipo": "tutti"}, **corpo})
        assert r.status_code == 403, (rotta, corpo, r.text)
        assert "non è scaricabile" in r.json()["detail"]


def test_dataset_inesistente(cliente):
    assert cliente.post("/api/storici/stima", json={"dataset": "pippo"}).status_code == 404


# ---------------------------------------------------------------------------------------------- formati
def test_csv_formato(cliente):
    r = _estrai(cliente, {"dataset": "meteo_giornaliero", "frequenza": "giornaliera", "dal": "2023-03-01",
                          "al": "2023-03-02", "variabili": ["t_media"], "area": {"tipo": "comuni", "comuni": ["037006"]}})
    assert r.headers["content-type"].startswith("text/csv")
    assert 'filename="eugenio_meteo_giornaliero_giornaliera_20230301-20230302.csv"' in r.headers["content-disposition"]
    righe = r.content.decode("utf-8-sig").split("\r\n")
    assert righe[0] == "istat;comune;provincia;data;t_media (°C)"
    assert righe[1].startswith("037006;Bologna;BO;2023-03-01;") and "," in righe[1].split(";")[-1]
    assert any(x.startswith("# Fonte: Weather data by Open-Meteo.com") for x in righe)
    assert any(x.startswith("# Licenza: CC BY 4.0") for x in righe)
    assert any(x.startswith("# Licenza EuGEnio: Elaborazioni e raccolte di EuGEnio: CC BY 4.0") and
               "citando EuGEnio come fonte" in x for x in righe)
    assert any(x.startswith("# Come citare: Fonte: EuGEnio – eugenio.statocritico.it (dati estratti il ")
               for x in righe)


def test_xlsx(cliente):
    from openpyxl import load_workbook
    corpo = {"dataset": "meteo_orario", "frequenza": "oraria", "dal": "2023-01-02", "al": "2023-01-02",
             "variabili": ["temperature_2m"], "area": {"tipo": "comuni", "comuni": ["048017"]}}
    r = _estrai(cliente, corpo, "xlsx")
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Dati", "Info"]
    righe = list(wb["Dati"].iter_rows(values_only=True))
    assert righe[0] == ("istat", "comune", "provincia", "ora_utc (UTC)", "ora_locale (ora italiana)", "temperatura (°C)")
    assert len(righe) == 25
    assert righe[1][4] == dt.datetime(2023, 1, 2, 0, 0) and righe[1][5] == pytest.approx(23 + 3)   # Firenze: k = 3
    info = {a: b for a, b, *_ in wb["Info"].iter_rows(values_only=True) if a}
    assert info["Dataset"].startswith("Meteo orario") and "Open-Meteo" in info["Fonte"]
    assert "CC BY 4.0" in info["Licenza"] and "Estratto il" in info and "Area" in info
    assert "CC BY 4.0" in info["Licenza EuGEnio"] and "citando EuGEnio come fonte" in info["Licenza EuGEnio"]
    assert info["Come citare"].startswith("Fonte: EuGEnio – eugenio.statocritico.it (dati estratti il ")
    gio = load_workbook(io.BytesIO(_estrai(cliente, {"dataset": "gradi_giorno", "frequenza": "giornaliera",
                                                     "dal": "2022-01-01", "al": "2022-01-03",
                                                     "area": {"tipo": "comuni", "comuni": ["033001"]}}, "xlsx").content))
    r1 = list(gio["Dati"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert r1[3] == dt.datetime(2022, 1, 1)   # data vera, non testo


def test_ods(cliente):
    import zipfile
    corpo = {"dataset": "meteo_orario", "frequenza": "oraria", "dal": "2023-01-02", "al": "2023-01-02",
             "variabili": ["temperature_2m"], "area": {"tipo": "comuni", "comuni": ["048017"]}}
    r = _estrai(cliente, corpo, "ods")
    assert r.headers["content-type"] == "application/vnd.oasis.opendocument.spreadsheet"
    assert 'filename="eugenio_meteo_orario_oraria_20230102-20230102.ods"' in r.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(r.content))
    primo = z.infolist()[0]
    assert primo.filename == "mimetype" and primo.compress_type == zipfile.ZIP_STORED
    assert z.read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"
    assert "META-INF/manifest.xml" in z.namelist()
    x = z.read("content.xml").decode("utf-8")
    assert "table:formula" not in x
    fogli = pd.read_excel(io.BytesIO(r.content), engine="odf", sheet_name=None)
    assert list(fogli) == ["Dati", "Info"]
    dati = fogli["Dati"]
    assert list(dati.columns) == ["istat", "comune", "provincia", "ora_utc (UTC)", "ora_locale (ora italiana)",
                                  "temperatura (°C)"]
    assert len(dati) == 24
    assert dati["temperatura (°C)"].iloc[0] == pytest.approx(23 + 3)               # numero, non testo
    assert pd.Timestamp(dati["ora_locale (ora italiana)"].iloc[0]) == pd.Timestamp(2023, 1, 2, 0, 0)  # data vera
    assert 'office:value-type="date" office:date-value="2023-01-02T00:00:00"' in x
    info = {a: b for a, b in fogli["Info"].iloc[:, :2].itertuples(index=False) if isinstance(a, str)}
    assert info["Dataset"].startswith("Meteo orario") and "Open-Meteo" in info["Fonte"]
    assert "CC BY 4.0" in info["Licenza EuGEnio"] and "citando EuGEnio come fonte" in info["Licenza EuGEnio"]
    assert info["Come citare"].startswith("Fonte: EuGEnio – eugenio.statocritico.it (dati estratti il ")
    stili = z.read("styles.xml").decode("utf-8")
    assert "Fonte: EuGEnio – eugenio.statocritico.it" in stili and "<text:page-number" in stili


def test_ods_testo_sicuro():
    """Caratteri speciali e di controllo nel testo: XML sempre valido."""
    import xml.dom.minidom
    cella = storici._cella_ods('a<b & "c"\x01')
    xml.dom.minidom.parseString(f'<r xmlns:table="t" xmlns:office="o" xmlns:text="x">{cella}</r>')
    assert "&lt;b &amp; &quot;c&quot;" in cella and "\x01" not in cella
    assert storici._cella_ods(float("nan")) == "<table:table-cell/>"
    assert 'office:value="3"' in storici._cella_ods(3) and "sì" in storici._cella_ods(True)


# ---------------------------------------------------------------------------------------------- DPR 412/93
def _dpr412_json() -> dict:
    return {"fonte": "DPR 412/93", "comuni": {
        "033001": [2715, "E", 179, "", ""],
        "033032": [2715, "E", 61, "", "ha incorporato Esempio (2800 GG, zona E)", [[2800, "E", 300, "ESEMPIO"]]],
        "037006": [None, "E", None, "A + B", "fusione di A, B", [[2259, "E", 54, "A"], [2300, "E", 80, "B"]]],
        "048017": [None, None, None, "", "non presente nel DPR"],
        "027042": [2345, "E", 1, "", ""],            # Veneto: il server non lo serve
    }}


def test_dpr412_dataset(cartella, tmp_path, monkeypatch):
    import json
    import shutil
    shutil.copy(cartella / "comuni.parquet", tmp_path / "comuni.parquet")
    monkeypatch.setattr(C, "DATI", tmp_path)
    monkeypatch.setattr(C, "REGIONI", ["08", "09"])
    monkeypatch.setattr(C, "SOLO_COMUNI", [])
    storici._comuni_cache.cache_clear()
    storici._disponibilita.cache_clear()
    storici._dpr412_cache.cache_clear()
    storici._storia.clear()
    app = FastAPI()
    app.include_router(storici.router)
    cl = TestClient(app)
    ids = [x["id"] for x in cl.get("/api/storici/catalogo").json()["dataset"]]
    assert "dpr412" not in ids                                   # senza file il dataset non compare
    (tmp_path / "dpr412.json").write_text(json.dumps(_dpr412_json()), encoding="utf-8")
    cat = {x["id"]: x for x in cl.get("/api/storici/catalogo").json()["dataset"]}
    assert "dpr412" in cat and "Normattiva" in cat["dpr412"]["fonte"] and "633/1941" in cat["dpr412"]["licenza"]
    r = _estrai(cl, {"dataset": "dpr412", "area": {"tipo": "tutti"}})
    df = _leggi_csv(r.content)
    assert set(df.istat) == {"033001", "033032", "037006", "048017"}       # Venezia fuori dalle regioni servite
    assert len(df) == 5                                                     # Bologna: due righe (fusione)
    bo = df[df.istat == "037006"]
    assert set(bo.tipo) == {"fusione"} and set(bo["gradi_giorno (°C·giorno)"]) == {2259, 2300}
    pc = df.set_index("istat").loc["033032"]
    assert pc.tipo == "incorporazione" and pc["gradi_giorno (°C·giorno)"] == 2715 and pc["altitudine (m)"] == 61
    assert df.set_index("istat").loc["048017"].tipo == "assente"
    testo = r.content.decode("utf-8-sig")
    assert "DPR 26 agosto 1993" in testo and "normattiva" in testo.lower()
    st = cl.post("/api/storici/stima", json={"dataset": "dpr412", "area": {"tipo": "provincia", "provincia": "PC"}})
    assert st.status_code == 200 and st.json()["righe"] == 2
    assert "2715;E;61" in testo and "048017;Firenze;FI;;;;" in testo                     # interi; vuoti se assente
    x = cl.post("/api/storici/estrai?formato=xlsx", json={"dataset": "dpr412", "area": {"tipo": "tutti"}})
    assert x.status_code == 200 and x.content[:2] == b"PK"
